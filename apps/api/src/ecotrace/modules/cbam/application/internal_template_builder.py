from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName

from ecotrace.modules.cbam.application.export_storage import (
    INTERNAL_MAPPING_VERSION,
    INTERNAL_TEMPLATE_CODE,
    INTERNAL_TEMPLATE_VERSION,
    sha256_bytes,
    template_storage_path,
)

SHEETS = (
    'Organization',
    'Installation',
    'Production',
    'Activities',
    'Purchased Inputs',
    'Allocation',
    'Factors',
    'Calculations',
    'Summary',
)

REPEATING_HEADERS: dict[str, tuple[str, ...]] = {
    'Production': ('product', 'quantity', 'unit', 'date', 'id'),
    'Activities': ('group', 'type', 'quantity', 'unit', 'data_source', 'date', 'id'),
    'Purchased Inputs': (
        'input_name',
        'supplier',
        'purchased_quantity',
        'consumed_quantity',
        'unit',
        'embedded_value',
        'embedded_unit',
        'id',
    ),
    'Allocation': (
        'source_type',
        'source_id',
        'method',
        'ratio',
        'allocated_quantity',
        'allocated_unit',
        'id',
    ),
    'Factors': (
        'source_type',
        'source_id',
        'definition',
        'selected_value',
        'unit',
        'precedence',
        'status',
        'id',
    ),
    'Calculations': (
        'source_type',
        'source_id',
        'quantity',
        'quantity_unit',
        'factor',
        'factor_unit',
        'result',
        'result_unit',
        'status',
        'error',
        'id',
    ),
}


def build_internal_template_bytes() -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for name in SHEETS:
        wb.create_sheet(name)

    banner = 'INTERNAL DEVELOPMENT TEMPLATE — NOT AN OFFICIAL CBAM SUBMISSION FORMAT'
    bold = Font(bold=True)

    org = wb['Organization']
    org['A1'] = banner
    org['A1'].font = bold
    org['A3'] = 'company_name'
    org['B3'] = None
    org['A4'] = 'organization_code'
    org['B4'] = None
    wb.defined_names.add(DefinedName(name='ORG_COMPANY_NAME', attr_text="'Organization'!$B$3"))
    wb.defined_names.add(DefinedName(name='ORG_CODE', attr_text="'Organization'!$B$4"))

    inst = wb['Installation']
    inst['A1'] = banner
    inst['A1'].font = bold
    inst['A3'] = 'installation_code'
    inst['B3'] = None
    inst['A4'] = 'installation_name'
    inst['B4'] = None
    inst['A5'] = 'reporting_period'
    inst['B5'] = None

    for sheet_name, headers in REPEATING_HEADERS.items():
        ws = wb[sheet_name]
        ws['A1'] = banner
        for idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=idx, value=header)
            cell.font = bold

    summary = wb['Summary']
    summary['A1'] = banner
    summary['A1'].font = bold
    summary['A3'] = 'label'
    summary['B3'] = 'value'
    summary['A3'].font = bold
    summary['B3'].font = bold
    labels = [
        ('production_record_count', 'B4'),
        ('activity_record_count', 'B5'),
        ('purchased_input_count', 'B6'),
        ('allocation_result_count', 'B7'),
        ('resolved_primary_factor_count', 'B8'),
        ('resolved_default_factor_count', 'B9'),
        ('unresolved_factor_count', 'B10'),
        ('ambiguous_factor_count', 'B11'),
        ('calculated_result_count', 'B12'),
        ('blocked_calculation_count', 'B13'),
        ('technical_total', 'B14'),
        ('technical_total_unit', 'B15'),
        ('export_readiness', 'B16'),
        ('export_date', 'B17'),
        ('disclaimer', 'B18'),
    ]
    for offset, (label, _cell) in enumerate(labels):
        row = 4 + offset
        summary.cell(row=row, column=1, value=label)
        summary.cell(row=row, column=2, value=None)
    summary['D3'] = 'formula_metric_row_count'
    summary['E3'] = '=COUNTA(A4:A18)'
    summary['A20'] = 'NOTICE'
    summary['B20'] = (
        'This workbook is an EcoTrace internal development template. '
        'It is not an official EU CBAM / SKDM submission format.'
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def write_internal_template_file() -> tuple[Path, str]:
    data = build_internal_template_bytes()
    path = template_storage_path(code=INTERNAL_TEMPLATE_CODE, version=INTERNAL_TEMPLATE_VERSION)
    path.write_bytes(data)
    return path, sha256_bytes(data)


def internal_mapping_specs() -> list[dict[str, object]]:
    specs: list[dict[str, object]] = [
        {
            'mapping_code': 'org.company_name',
            'source_type': 'ORGANIZATION',
            'source_path': 'organization.name',
            'worksheet_name': 'Organization',
            'destination_type': 'NAMED_RANGE',
            'destination_reference': 'ORG_COMPANY_NAME',
            'value_type': 'STRING',
            'required': True,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'org.code',
            'source_type': 'ORGANIZATION',
            'source_path': 'organization.code',
            'worksheet_name': 'Organization',
            'destination_type': 'NAMED_RANGE',
            'destination_reference': 'ORG_CODE',
            'value_type': 'STRING',
            'required': False,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'inst.code',
            'source_type': 'INSTALLATION',
            'source_path': 'installation.code',
            'worksheet_name': 'Installation',
            'destination_type': 'CELL',
            'destination_reference': 'B3',
            'value_type': 'STRING',
            'required': True,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'inst.name',
            'source_type': 'INSTALLATION',
            'source_path': 'installation.name',
            'worksheet_name': 'Installation',
            'destination_type': 'CELL',
            'destination_reference': 'B4',
            'value_type': 'STRING',
            'required': True,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'period.label',
            'source_type': 'REPORTING_PERIOD',
            'source_path': 'reporting_period.label',
            'worksheet_name': 'Installation',
            'destination_type': 'CELL',
            'destination_reference': 'B5',
            'value_type': 'STRING',
            'required': True,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'summary.disclaimer',
            'source_type': 'CONSTANT',
            'source_path': 'constant.internal_disclaimer',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B18',
            'value_type': 'STRING',
            'required': True,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'summary.export_date',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.export_date',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B17',
            'value_type': 'STRING',
            'required': True,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'summary.readiness',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.export_readiness',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B16',
            'value_type': 'STRING',
            'required': True,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'summary.production_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.production_record_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B4',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.activity_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.activity_record_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B5',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.purchased_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.purchased_input_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B6',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.allocation_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.allocation_result_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B7',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.primary_factor_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.resolved_primary_factor_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B8',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.default_factor_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.resolved_default_factor_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B9',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.unresolved_factor_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.unresolved_factor_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B10',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.ambiguous_factor_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.ambiguous_factor_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B11',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.calculated_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.calculated_result_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B12',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.blocked_count',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.blocked_calculation_count',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B13',
            'value_type': 'NUMBER',
            'required': True,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.technical_total',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.technical_total',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B14',
            'value_type': 'NUMBER',
            'required': False,
            'transformation_code': 'DECIMAL_TO_NUMBER',
        },
        {
            'mapping_code': 'summary.technical_total_unit',
            'source_type': 'CALCULATED_SUMMARY',
            'source_path': 'summary.technical_total_unit',
            'worksheet_name': 'Summary',
            'destination_type': 'CELL',
            'destination_reference': 'B15',
            'value_type': 'UNIT',
            'required': False,
            'transformation_code': 'UNIT_DISPLAY',
        },
        {
            'mapping_code': 'rows.production',
            'source_type': 'PRODUCTION_RECORD',
            'source_path': 'production.rows',
            'worksheet_name': 'Production',
            'destination_type': 'REPEATING_ROW',
            'destination_reference': '3|product,quantity,unit,date,id',
            'value_type': 'STRING',
            'required': False,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'rows.activities',
            'source_type': 'ACTIVITY_RECORD',
            'source_path': 'activity.rows',
            'worksheet_name': 'Activities',
            'destination_type': 'REPEATING_ROW',
            'destination_reference': '3|group,type,quantity,unit,data_source,date,id',
            'value_type': 'STRING',
            'required': False,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'rows.purchased',
            'source_type': 'PURCHASED_INPUT',
            'source_path': 'purchased.rows',
            'worksheet_name': 'Purchased Inputs',
            'destination_type': 'REPEATING_ROW',
            'destination_reference': (
                '3|input_name,supplier,purchased_quantity,consumed_quantity,'
                'unit,embedded_value,embedded_unit,id'
            ),
            'value_type': 'STRING',
            'required': False,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'rows.allocation',
            'source_type': 'ALLOCATION_RESULT',
            'source_path': 'allocation.rows',
            'worksheet_name': 'Allocation',
            'destination_type': 'REPEATING_ROW',
            'destination_reference': (
                '3|source_type,source_id,method,ratio,allocated_quantity,allocated_unit,id'
            ),
            'value_type': 'STRING',
            'required': False,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'rows.factors',
            'source_type': 'FACTOR_RESOLUTION',
            'source_path': 'factor.rows',
            'worksheet_name': 'Factors',
            'destination_type': 'REPEATING_ROW',
            'destination_reference': (
                '3|source_type,source_id,definition,selected_value,unit,precedence,status,id'
            ),
            'value_type': 'STRING',
            'required': False,
            'transformation_code': 'NONE',
        },
        {
            'mapping_code': 'rows.calculations',
            'source_type': 'CALCULATION_RESULT',
            'source_path': 'calculation.rows',
            'worksheet_name': 'Calculations',
            'destination_type': 'REPEATING_ROW',
            'destination_reference': (
                '3|source_type,source_id,quantity,quantity_unit,factor,factor_unit,'
                'result,result_unit,status,error,id'
            ),
            'value_type': 'STRING',
            'required': False,
            'transformation_code': 'NONE',
        },
    ]
    for spec in specs:
        spec['notes'] = f'{INTERNAL_MAPPING_VERSION}; internal development mapping only'
    return specs
