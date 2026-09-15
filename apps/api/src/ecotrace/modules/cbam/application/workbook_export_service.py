from __future__ import annotations

import json
import shutil
import uuid
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ecotrace.core.config import get_settings
from ecotrace.core.exceptions import NotFoundError, ValidationAppError
from ecotrace.modules.cbam.application.collection_guards import get_binding_for_org
from ecotrace.modules.cbam.application.export_context import (
    ExportDataContext,
    load_export_context,
    repeating_rows,
    scalar_source_value,
)
from ecotrace.modules.cbam.application.export_readiness_service import assess_export_readiness
from ecotrace.modules.cbam.application.export_storage import (
    export_run_dir,
    relative_uri,
    resolve_uri,
    sha256_file,
)
from ecotrace.modules.cbam.application.export_template_service import (
    get_active_template_for_org,
    list_template_mappings,
    mapping_checksum,
    verify_template_file,
)
from ecotrace.modules.cbam.application.permissions import require_cbam_configure, require_cbam_view
from ecotrace.modules.cbam.infrastructure.models import (
    CbamExportArtifact,
    CbamExportMapping,
    CbamExportRun,
)
from ecotrace.modules.identity.infrastructure.models import User
from ecotrace.shared.application.audit import write_audit_log
from ecotrace.shared.domain.schemas import CamelModel, Page, paginate

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FORMULA_INJECTION_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


def _excel_safe_cell_value(value: object | None) -> object | None:
    if not isinstance(value, str):
        return value
    if value.startswith(_FORMULA_INJECTION_PREFIXES):
        return f"'{value}"
    return value


class ExportCreateRequest(CamelModel):
    export_template_id: uuid.UUID | None = None
    calculation_run_id: uuid.UUID | None = None


class ExportRunResponse(CamelModel):
    id: uuid.UUID
    organization_id: uuid.UUID
    reporting_period_binding_id: uuid.UUID
    export_template_id: uuid.UUID
    calculation_run_id: uuid.UUID | None
    status: str
    template_version: str
    mapping_version: str
    mapping_checksum: str
    template_checksum: str
    input_checksum: str | None
    started_at: datetime | None
    completed_at: datetime | None
    error_message: str | None
    warning_summary: str | None
    application_version: str | None


class ExportArtifactResponse(CamelModel):
    id: uuid.UUID
    organization_id: uuid.UUID
    export_run_id: uuid.UUID
    artifact_type: str
    file_name: str
    storage_uri: str
    mime_type: str
    file_size_bytes: int
    sha256: str
    created_at: datetime


def _run_response(row: CbamExportRun) -> ExportRunResponse:
    return ExportRunResponse.model_validate(row)


def _artifact_response(row: CbamExportArtifact) -> ExportArtifactResponse:
    return ExportArtifactResponse.model_validate(row)


def _collect_formula_cells(wb: Workbook) -> dict[str, str]:
    formulas: dict[str, str] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas[f"{ws.title}!{cell.coordinate}"] = value
    return formulas


def _transform(value: object | None, transformation: str | None) -> object | None:
    if value is None:
        return None
    code = transformation or "NONE"
    if code in {"NONE", "UNIT_DISPLAY", "ENUM_TO_DISPLAY_LABEL"}:
        return value if not isinstance(value, Decimal) else float(value)
    if code == "DECIMAL_TO_NUMBER":
        if isinstance(value, Decimal):
            return float(value)
        if isinstance(value, (int, float)):
            return value
        return float(str(value))
    if code == "DATE_TO_EXCEL_DATE":
        if isinstance(value, datetime):
            return value.date()
        return value
    if code == "DATETIME_TO_EXCEL_DATETIME":
        return value
    if code == "BOOLEAN_TO_YES_NO":
        return "YES" if bool(value) else "NO"
    raise ValidationAppError(f"Unsupported transformation: {code}")


def _write_cell(wb: Workbook, worksheet_name: str, reference: str, value: object | None) -> None:
    if worksheet_name not in wb.sheetnames:
        raise ValidationAppError(f"Worksheet missing: {worksheet_name}")
    ws = wb[worksheet_name]
    cell = ws[reference]
    if isinstance(cell.value, str) and cell.value.startswith("="):
        raise ValidationAppError(f"Refusing to overwrite formula cell {worksheet_name}!{reference}")
    cell.value = _excel_safe_cell_value(value)


def _write_named_range(wb: Workbook, name: str, value: object | None) -> None:
    defined = wb.defined_names.get(name)
    if defined is None:
        raise ValidationAppError(f"Named range missing: {name}")
    destinations = list(defined.destinations)
    if not destinations:
        raise ValidationAppError(f"Named range has no destination: {name}")
    sheet_title, coord = destinations[0]
    _write_cell(wb, sheet_title, coord.replace("$", ""), value)


def _write_repeating(
    wb: Workbook,
    mapping: CbamExportMapping,
    rows: list[dict[str, object | None]],
) -> None:
    if "|" not in mapping.destination_reference:
        raise ValidationAppError(f"Invalid repeating destination for {mapping.mapping_code}")
    start_raw, cols_raw = mapping.destination_reference.split("|", 1)
    start_row = int(start_raw)
    columns = [c.strip() for c in cols_raw.split(",") if c.strip()]
    ws = wb[mapping.worksheet_name]
    for offset, row in enumerate(rows):
        excel_row = start_row + offset
        for col_idx, key in enumerate(columns, start=1):
            cell = ws.cell(row=excel_row, column=col_idx)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                raise ValidationAppError(
                    f"Refusing to overwrite formula at {mapping.worksheet_name}!{cell.coordinate}"
                )
            value = row.get(key)
            if isinstance(value, Decimal):
                cell.value = float(value)
            elif isinstance(value, (date, datetime)):
                cell.value = value.isoformat() if not isinstance(value, datetime) else value
            else:
                cell.value = _excel_safe_cell_value(value)


def _populate_workbook(
    wb: Workbook,
    mappings: list[CbamExportMapping],
    ctx: ExportDataContext,
) -> list[str]:
    mapped_fields: list[str] = []
    for mapping in mappings:
        if mapping.destination_type == "REPEATING_ROW":
            rows = repeating_rows(ctx, mapping.source_path)
            _write_repeating(wb, mapping, rows)
            mapped_fields.append(mapping.mapping_code)
            continue
        value = scalar_source_value(ctx, mapping.source_path)
        transformed = _transform(value, mapping.transformation_code)
        if mapping.required and transformed is None:
            raise ValidationAppError(f"Required mapping missing value: {mapping.mapping_code}")
        if mapping.destination_type == "CELL":
            _write_cell(wb, mapping.worksheet_name, mapping.destination_reference, transformed)
        elif mapping.destination_type == "NAMED_RANGE":
            _write_named_range(wb, mapping.destination_reference, transformed)
        else:
            raise ValidationAppError(
                f"Unsupported destination type for Phase 6: {mapping.destination_type}"
            )
        mapped_fields.append(mapping.mapping_code)
    return mapped_fields


def create_export_run(
    db: Session,
    user: User,
    organization_id: uuid.UUID,
    binding_id: uuid.UUID,
    payload: ExportCreateRequest | None = None,
    *,
    request_id: str | None = None,
    ip_address: str | None = None,
    user_agent: str | None = None,
) -> ExportRunResponse:
    require_cbam_configure(db, user, organization_id)
    get_binding_for_org(db, organization_id, binding_id)
    payload = payload or ExportCreateRequest()

    readiness = assess_export_readiness(
        db,
        user,
        organization_id,
        binding_id,
        template_id=payload.export_template_id,
        calculation_run_id=payload.calculation_run_id,
    )
    if readiness.status == "NOT_READY":
        raise ValidationAppError(
            "Export is not ready.",
            details=[{"field": "readiness", "message": "; ".join(readiness.blocking_issues)}],
        )

    template = get_active_template_for_org(
        db, organization_id, payload.export_template_id or readiness.suggested_template_id
    )
    if template.template_type == "OFFICIAL_CBAM_TEMPLATE":
        raise ValidationAppError(
            "Official CBAM workbook mapping is BLOCKED pending domain-expert/template delivery."
        )

    template_path = verify_template_file(template)
    mappings = list_template_mappings(db, template.id)
    map_checksum = mapping_checksum(mappings)
    calc_run_id = payload.calculation_run_id or readiness.calculation_run_id

    run = CbamExportRun(
        organization_id=organization_id,
        reporting_period_binding_id=binding_id,
        export_template_id=template.id,
        calculation_run_id=calc_run_id,
        status="RUNNING",
        template_version=template.version,
        mapping_version=template.mapping_version,
        mapping_checksum=map_checksum,
        template_checksum=template.checksum,
        started_at=datetime.now(UTC),
        application_version=get_settings().app_version,
        created_by_user_id=user.id,
        warning_summary="; ".join(readiness.warnings) if readiness.warnings else None,
    )
    db.add(run)
    db.flush()

    out_dir = export_run_dir(
        organization_id=organization_id,
        binding_id=binding_id,
        export_run_id=run.id,
    )
    output_xlsx = out_dir / f"skdm-export-{run.id}.xlsx"
    try:
        shutil.copy2(template_path, output_xlsx)
        before_wb = load_workbook(template_path)
        before_formulas = _collect_formula_cells(before_wb)
        before_count = len(before_formulas)

        wb = load_workbook(output_xlsx)
        wb.properties.creator = "EcoTrace AI CBAM Export"
        wb.properties.lastModifiedBy = "EcoTrace AI CBAM Export"

        ctx = load_export_context(
            db,
            organization_id=organization_id,
            binding_id=binding_id,
            calculation_run_id=calc_run_id,
        )
        ctx.summary_metrics["export_readiness"] = readiness.status
        mapped_fields = _populate_workbook(wb, mappings, ctx)
        wb.save(output_xlsx)

        after_wb = load_workbook(output_xlsx)
        after_formulas = _collect_formula_cells(after_wb)
        if len(after_formulas) != before_count:
            raise ValidationAppError(
                f"Formula preservation failed: before={before_count} after={len(after_formulas)}"
            )
        for key, formula in before_formulas.items():
            if after_formulas.get(key) != formula:
                raise ValidationAppError(f"Formula cell changed unexpectedly: {key}")

        xlsx_checksum = sha256_file(output_xlsx)
        run.input_checksum = xlsx_checksum
        artifact_xlsx = CbamExportArtifact(
            organization_id=organization_id,
            export_run_id=run.id,
            artifact_type="XLSX",
            file_name=output_xlsx.name,
            storage_uri=relative_uri(output_xlsx),
            mime_type=XLSX_MIME,
            file_size_bytes=output_xlsx.stat().st_size,
            sha256=xlsx_checksum,
        )
        db.add(artifact_xlsx)

        manifest = {
            "exportRunId": str(run.id),
            "organizationId": str(organization_id),
            "reportingPeriodBindingId": str(binding_id),
            "templateCode": template.code,
            "templateVersion": template.version,
            "templateChecksum": template.checksum,
            "mappingVersion": template.mapping_version,
            "mappingChecksum": map_checksum,
            "calculationRunId": str(calc_run_id) if calc_run_id else None,
            "generatedAt": datetime.now(UTC).isoformat(),
            "applicationVersion": get_settings().app_version,
            "mappedFields": mapped_fields,
            "warnings": readiness.warnings,
            "artifactChecksum": xlsx_checksum,
            "traceability": ctx.mapped_trace_ids,
            "officialMappingBlocked": True,
            "disclaimer": (
                "INTERNAL DEVELOPMENT TEMPLATE — NOT AN OFFICIAL CBAM SUBMISSION FORMAT"
            ),
        }
        manifest_path = out_dir / "export-manifest.json"
        manifest_bytes = json.dumps(manifest, indent=2, sort_keys=True).encode("utf-8")
        manifest_path.write_bytes(manifest_bytes)
        db.add(
            CbamExportArtifact(
                organization_id=organization_id,
                export_run_id=run.id,
                artifact_type="JSON_SUMMARY",
                file_name=manifest_path.name,
                storage_uri=relative_uri(manifest_path),
                mime_type="application/json",
                file_size_bytes=len(manifest_bytes),
                sha256=sha256_file(manifest_path),
            )
        )

        from ecotrace.modules.cbam.application.period_summary_service import (
            render_summary_html,
            render_summary_json,
        )

        summary = render_summary_json(ctx, readiness.status)
        summary_path = out_dir / "skdm-period-summary.json"
        summary_bytes = json.dumps(summary, indent=2, sort_keys=True, default=str).encode("utf-8")
        summary_path.write_bytes(summary_bytes)
        db.add(
            CbamExportArtifact(
                organization_id=organization_id,
                export_run_id=run.id,
                artifact_type="JSON_SUMMARY",
                file_name=summary_path.name,
                storage_uri=relative_uri(summary_path),
                mime_type="application/json",
                file_size_bytes=len(summary_bytes),
                sha256=sha256_file(summary_path),
            )
        )
        html = render_summary_html(summary)
        html_path = out_dir / "skdm-period-summary.html"
        html_bytes = html.encode("utf-8")
        html_path.write_bytes(html_bytes)
        db.add(
            CbamExportArtifact(
                organization_id=organization_id,
                export_run_id=run.id,
                artifact_type="HTML_REPORT",
                file_name=html_path.name,
                storage_uri=relative_uri(html_path),
                mime_type="text/html",
                file_size_bytes=len(html_bytes),
                sha256=sha256_file(html_path),
            )
        )

        run.status = (
            "COMPLETED_WITH_WARNINGS" if readiness.status == "READY_WITH_WARNINGS" else "COMPLETED"
        )
        run.completed_at = datetime.now(UTC)
        write_audit_log(
            db,
            action="cbam.export.generated",
            actor_user_id=user.id,
            organization_id=organization_id,
            entity_type="cbam_export_run",
            entity_id=str(run.id),
            metadata={
                "reportingPeriodBindingId": str(binding_id),
                "templateId": str(template.id),
                "templateVersion": template.version,
                "calculationRunId": str(calc_run_id) if calc_run_id else None,
                "artifactChecksum": xlsx_checksum,
                "status": run.status,
            },
            request_id=request_id,
            ip_address=ip_address,
            user_agent=user_agent,
        )
        db.commit()
        db.refresh(run)
        return _run_response(run)
    except Exception as exc:
        run.status = "FAILED"
        run.error_message = str(exc)
        run.completed_at = datetime.now(UTC)
        write_audit_log(
            db,
            action="cbam.export.failed",
            actor_user_id=user.id,
            organization_id=organization_id,
            entity_type="cbam_export_run",
            entity_id=str(run.id),
            metadata={
                "reportingPeriodBindingId": str(binding_id),
                "templateId": str(template.id),
                "error": str(exc),
            },
            request_id=request_id,
            ip_address=ip_address,
            user_agent=user_agent,
        )
        db.commit()
        raise


def list_export_runs(
    db: Session,
    user: User,
    organization_id: uuid.UUID,
    binding_id: uuid.UUID,
    *,
    page: int,
    page_size: int,
) -> Page[ExportRunResponse]:
    require_cbam_view(db, user, organization_id)
    get_binding_for_org(db, organization_id, binding_id)
    stmt = select(CbamExportRun).where(
        CbamExportRun.organization_id == organization_id,
        CbamExportRun.reporting_period_binding_id == binding_id,
    )
    total = db.execute(select(func.count()).select_from(stmt.subquery())).scalar_one()
    rows = list(
        db.execute(
            stmt.order_by(CbamExportRun.created_at.desc())
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
        .scalars()
        .all()
    )
    return paginate(
        [_run_response(r) for r in rows],
        page=page,
        page_size=page_size,
        total_items=int(total),
    )


def get_export_run(
    db: Session,
    user: User,
    organization_id: uuid.UUID,
    export_run_id: uuid.UUID,
) -> ExportRunResponse:
    require_cbam_view(db, user, organization_id)
    row = db.execute(
        select(CbamExportRun).where(
            CbamExportRun.id == export_run_id,
            CbamExportRun.organization_id == organization_id,
        )
    ).scalar_one_or_none()
    if row is None:
        raise NotFoundError("Export run not found.")
    return _run_response(row)


def list_export_artifacts(
    db: Session,
    user: User,
    organization_id: uuid.UUID,
    export_run_id: uuid.UUID,
) -> list[ExportArtifactResponse]:
    require_cbam_view(db, user, organization_id)
    get_export_run(db, user, organization_id, export_run_id)
    rows = list(
        db.execute(
            select(CbamExportArtifact)
            .where(
                CbamExportArtifact.organization_id == organization_id,
                CbamExportArtifact.export_run_id == export_run_id,
            )
            .order_by(CbamExportArtifact.created_at.asc())
        ).scalars()
    )
    return [_artifact_response(r) for r in rows]


def download_export_artifact(
    db: Session,
    user: User,
    organization_id: uuid.UUID,
    artifact_id: uuid.UUID,
    *,
    request_id: str | None = None,
    ip_address: str | None = None,
    user_agent: str | None = None,
) -> tuple[CbamExportArtifact, Path]:
    require_cbam_view(db, user, organization_id)
    row = db.execute(
        select(CbamExportArtifact).where(
            CbamExportArtifact.id == artifact_id,
            CbamExportArtifact.organization_id == organization_id,
        )
    ).scalar_one_or_none()
    if row is None:
        raise NotFoundError("Export artifact not found.")
    path = resolve_uri(row.storage_uri)
    if not path.is_file():
        raise NotFoundError("Export artifact file not found.")
    write_audit_log(
        db,
        action="cbam.export.downloaded",
        actor_user_id=user.id,
        organization_id=organization_id,
        entity_type="cbam_export_artifact",
        entity_id=str(row.id),
        metadata={
            "exportRunId": str(row.export_run_id),
            "artifactType": row.artifact_type,
            "checksum": row.sha256,
        },
        request_id=request_id,
        ip_address=ip_address,
        user_agent=user_agent,
    )
    db.commit()
    return row, path
