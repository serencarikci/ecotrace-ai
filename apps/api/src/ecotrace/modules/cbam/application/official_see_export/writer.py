"""Copy official template → clear example → write INPUT via surgical ZIP/XML patching."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ecotrace.core.exceptions import BusinessRuleError
from ecotrace.modules.cbam.application.official_see_export.constants import (
    CODE_EXAMPLE_LEAKAGE,
    CODE_FORMULA_PRESERVATION_FAILED,
)
from ecotrace.modules.cbam.application.official_see_export.context import (
    OfficialSeeExportContext,
    resolve_source,
)
from ecotrace.modules.cbam.application.official_see_export.forensics import collect_formula_map
from ecotrace.modules.cbam.application.official_see_export.leakage import scan_example_leakage
from ecotrace.modules.cbam.application.official_see_export.mapping.loader import (
    MappingManifest,
    load_manifest,
)
from ecotrace.modules.cbam.application.official_see_export.package_writer import (
    apply_cell_patches,
)


def build_used_input_map(
    ctx: OfficialSeeExportContext,
    manifest: MappingManifest,
) -> dict[tuple[str, str], Any]:
    used: dict[tuple[str, str], Any] = {}
    for entry in manifest.by_direction('INPUT'):
        value = resolve_source(ctx, entry.source)
        if value is None:
            continue
        used[(entry.sheet, entry.cell)] = value
    return used


def build_cell_patches(
    ctx: OfficialSeeExportContext,
    manifest: MappingManifest,
) -> tuple[dict[tuple[str, str], Any | None], set[tuple[str, str]], dict[tuple[str, str], Any]]:
    """Plan CLEAR_EXAMPLE + unused INPUT clears and INPUT writes.

    Returns (patches, used_keys, used_values).
    None values in patches mean clear.
    """
    used = build_used_input_map(ctx, manifest)
    used_keys = set(used.keys())
    patches: dict[tuple[str, str], Any | None] = {}

    for entry in manifest.by_direction('CLEAR_EXAMPLE'):
        patches[(entry.sheet, entry.cell)] = None

    for entry in manifest.by_direction('INPUT'):
        key = (entry.sheet, entry.cell)
        if key in used_keys:
            patches[key] = used[key]
        else:
            # Unused mapped INPUT slots must be empty.
            patches[key] = None

    return patches, used_keys, used


def write_official_see_workbook(
    *,
    template_path: Path,
    output_path: Path,
    ctx: OfficialSeeExportContext,
    manifest: MappingManifest | None = None,
) -> dict[str, Any]:
    """Copy official template, clear examples, write EcoTrace INPUT values, preserve formulas.

    Uses surgical ZIP/XML patching (not openpyxl.save) so conditionalFormatting
    extension lists and other unsupported OOXML parts survive.
    """
    manifest = manifest or load_manifest()
    before_wb = load_workbook(template_path, data_only=False)
    before_formulas = collect_formula_map(before_wb)
    before_count = len(before_formulas)

    patches, used_keys, used = build_cell_patches(ctx, manifest)
    write_meta = apply_cell_patches(
        template_path=template_path,
        output_path=output_path,
        patches=patches,
        set_full_calc_on_load=True,
    )

    after_wb = load_workbook(output_path, data_only=False)
    leaks = scan_example_leakage(after_wb, manifest, used_input_keys=used_keys)
    if leaks:
        raise BusinessRuleError(
            'Example data leakage detected after clear/write.',
            code=CODE_EXAMPLE_LEAKAGE,
            details=[
                {
                    'code': CODE_EXAMPLE_LEAKAGE,
                    'kind': f.kind,
                    'sheet': f.sheet,
                    'cell': f.cell,
                    'detail': f.detail,
                }
                for f in leaks[:50]
            ],
        )

    after_formulas = collect_formula_map(after_wb)
    if len(after_formulas) != before_count:
        raise BusinessRuleError(
            f'Formula preservation failed: before={before_count} after={len(after_formulas)}',
            code=CODE_FORMULA_PRESERVATION_FAILED,
            details=[{'code': CODE_FORMULA_PRESERVATION_FAILED}],
        )
    samples_checked = 0
    for entry in manifest.formula_entries()[:200]:
        key = f'{entry.sheet}!{entry.cell}'
        actual = after_formulas.get(key)
        if entry.expected_formula and actual != entry.expected_formula:
            raise BusinessRuleError(
                f'Formula cell changed unexpectedly: {key}',
                code=CODE_FORMULA_PRESERVATION_FAILED,
                details=[
                    {
                        'code': CODE_FORMULA_PRESERVATION_FAILED,
                        'sheet': entry.sheet,
                        'cell': entry.cell,
                        'expected': entry.expected_formula,
                        'actual': actual,
                    }
                ],
            )
        samples_checked += 1
    for key, formula in before_formulas.items():
        if after_formulas.get(key) != formula:
            raise BusinessRuleError(
                f'Formula cell changed unexpectedly: {key}',
                code=CODE_FORMULA_PRESERVATION_FAILED,
                details=[{'code': CODE_FORMULA_PRESERVATION_FAILED, 'cell': key}],
            )

    return {
        'clearedCells': write_meta['clearedCells'],
        'writtenInputs': len(used),
        'patchedCells': write_meta['clearedCells'] + write_meta['writtenInputs'],
        'formulaCount': before_count,
        'formulaSamplesChecked': samples_checked,
        'modifiedParts': write_meta['modifiedParts'],
    }
