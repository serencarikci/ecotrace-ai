"""Workbook forensics helpers for Official SEE template inspection (tests/docs)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook


@dataclass(frozen=True, slots=True)
class SheetForensics:
    title: str
    state: str
    formula_count: int
    data_validation_count: int


@dataclass(frozen=True, slots=True)
class WorkbookForensics:
    sheet_names: tuple[str, ...]
    sheets: tuple[SheetForensics, ...]
    named_range_count: int
    named_ranges: tuple[str, ...]
    total_formula_count: int


def collect_formula_map(wb: Workbook) -> dict[str, str]:
    formulas: dict[str, str] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas[f"{ws.title}!{cell.coordinate}"] = value
    return formulas


def inspect_workbook(path: Path) -> WorkbookForensics:
    wb = load_workbook(path, data_only=False)
    sheets: list[SheetForensics] = []
    total = 0
    for ws in wb.worksheets:
        fcount = 0
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    fcount += 1
        total += fcount
        dv = getattr(ws, "data_validations", None)
        dv_count = len(dv.dataValidation) if dv is not None else 0
        sheets.append(
            SheetForensics(
                title=ws.title,
                state=str(ws.sheet_state),
                formula_count=fcount,
                data_validation_count=dv_count,
            )
        )
    names = tuple(sorted(wb.defined_names.keys()))
    return WorkbookForensics(
        sheet_names=tuple(wb.sheetnames),
        sheets=tuple(sheets),
        named_range_count=len(names),
        named_ranges=names,
        total_formula_count=total,
    )


def cell_value(wb: Workbook, sheet: str, cell: str) -> Any:
    return wb[sheet][cell].value


def scan_string_values(wb: Workbook, needles: tuple[str, ...]) -> list[dict[str, str]]:
    hits: list[dict[str, str]] = []
    lower_needles = tuple(n.lower() for n in needles)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or value.startswith("="):
                    continue
                low = value.lower()
                for needle, needle_low in zip(needles, lower_needles, strict=True):
                    if needle_low in low or value == needle:
                        hits.append(
                            {
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "value": value,
                                "needle": needle,
                            }
                        )
                        break
    return hits
