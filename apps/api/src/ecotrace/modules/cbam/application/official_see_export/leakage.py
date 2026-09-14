"""Detect leftover example identifiers and unused mapped slots with values."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from ecotrace.modules.cbam.application.official_see_export.constants import (
    EXAMPLE_IDENTIFIERS,
)
from ecotrace.modules.cbam.application.official_see_export.mapping.loader import MappingManifest

# Only EcoTrace-owned data entry sheets are scanned for example residue.
_DATA_SHEETS = frozenset(
    {
        'A_InstData',
        'B_EmInst',
        'C_Emissions&Energy',
        'D_Processes',
        'E_PurchPrec',
        'Summary_Products',
    }
)


@dataclass(frozen=True, slots=True)
class LeakageFinding:
    kind: str
    sheet: str
    cell: str
    detail: str


def scan_example_leakage(
    wb: Workbook | Path | str,
    manifest: MappingManifest,
    *,
    used_input_keys: set[tuple[str, str]],
    identifiers: tuple[str, ...] = EXAMPLE_IDENTIFIERS,
) -> list[LeakageFinding]:
    workbook = (
        load_workbook(Path(wb), data_only=False) if isinstance(wb, (str, Path)) else wb
    )

    findings: list[LeakageFinding] = []
    lower_needles = tuple(n.lower() for n in identifiers)

    watch_keys = {
        (e.sheet, e.cell)
        for e in manifest.entries
        if e.direction in {'CLEAR_EXAMPLE', 'INPUT'} and e.sheet in _DATA_SHEETS
    }
    for sheet, cell in watch_keys:
        if sheet not in workbook.sheetnames:
            continue
        value = workbook[sheet][cell].value
        if not isinstance(value, str) or value.startswith('='):
            continue
        low = value.lower().strip()
        for needle, needle_low in zip(identifiers, lower_needles, strict=True):
            # Exact (case-insensitive) match only — substring matching false-positives
            # legitimate CBAM product names that contain words like "Screws" / "nuts".
            if low == needle_low:
                findings.append(
                    LeakageFinding(
                        kind='EXAMPLE_IDENTIFIER',
                        sheet=sheet,
                        cell=cell,
                        detail=f'{needle!r} in {value!r}',
                    )
                )
                break

    for entry in manifest.by_direction('INPUT'):
        if entry.sheet not in _DATA_SHEETS:
            continue
        key = (entry.sheet, entry.cell)
        if key in used_input_keys:
            continue
        if entry.sheet not in workbook.sheetnames:
            continue
        value = workbook[entry.sheet][entry.cell].value
        if value is None:
            continue
        if isinstance(value, str) and value.startswith('='):
            continue
        findings.append(
            LeakageFinding(
                kind='UNUSED_SLOT_VALUE',
                sheet=entry.sheet,
                cell=entry.cell,
                detail=repr(value)[:200],
            )
        )

    return findings
