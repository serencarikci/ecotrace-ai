from __future__ import annotations

import hashlib
import json
import uuid
from datetime import UTC, datetime
from pathlib import Path

from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from ecotrace.core.exceptions import NotFoundError, ValidationAppError
from ecotrace.modules.cbam.application.export_storage import (
    INTERNAL_MAPPING_VERSION,
    INTERNAL_TEMPLATE_CODE,
    INTERNAL_TEMPLATE_VERSION,
    relative_uri,
    resolve_uri,
    sha256_file,
)
from ecotrace.modules.cbam.application.internal_template_builder import (
    SHEETS,
    internal_mapping_specs,
    write_internal_template_file,
)
from ecotrace.modules.cbam.application.permissions import require_cbam_view
from ecotrace.modules.cbam.infrastructure.models import CbamExportMapping, CbamExportTemplate
from ecotrace.modules.identity.infrastructure.models import User
from ecotrace.shared.domain.schemas import CamelModel, Page, paginate


class ExportTemplateResponse(CamelModel):
    id: uuid.UUID
    organization_id: uuid.UUID | None
    code: str
    name: str
    template_type: str
    version: str
    mapping_version: str
    storage_uri: str
    checksum: str
    status: str
    description: str | None
    activated_at: datetime | None
    archived_at: datetime | None
    official_mapping_blocked: bool = True


def _to_template(row: CbamExportTemplate) -> ExportTemplateResponse:
    return ExportTemplateResponse(
        id=row.id,
        organization_id=row.organization_id,
        code=row.code,
        name=row.name,
        template_type=row.template_type,
        version=row.version,
        mapping_version=row.mapping_version,
        storage_uri=row.storage_uri,
        checksum=row.checksum,
        status=row.status,
        description=row.description,
        activated_at=row.activated_at,
        archived_at=row.archived_at,
        official_mapping_blocked=row.template_type != "INTERNAL_SKDM",
    )


def mapping_checksum(mappings: list[CbamExportMapping]) -> str:
    payload = [
        {
            "mappingCode": m.mapping_code,
            "sourceType": m.source_type,
            "sourcePath": m.source_path,
            "worksheet": m.worksheet_name,
            "destinationType": m.destination_type,
            "destinationReference": m.destination_reference,
            "valueType": m.value_type,
            "required": m.required,
            "transformation": m.transformation_code,
        }
        for m in sorted(mappings, key=lambda x: x.mapping_code)
    ]
    raw = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def ensure_internal_export_template(db: Session) -> CbamExportTemplate:
    existing = db.execute(
        select(CbamExportTemplate).where(
            CbamExportTemplate.organization_id.is_(None),
            CbamExportTemplate.code == INTERNAL_TEMPLATE_CODE,
            CbamExportTemplate.version == INTERNAL_TEMPLATE_VERSION,
        )
    ).scalar_one_or_none()
    if existing is not None:
        path = resolve_uri(existing.storage_uri)
        if not path.is_file() or sha256_file(path) != existing.checksum:
            path, checksum = write_internal_template_file()
            existing.storage_uri = relative_uri(path)
            existing.checksum = checksum
        if existing.status == "DRAFT":
            existing.status = "ACTIVE"
            existing.activated_at = datetime.now(UTC)
    else:
        path, checksum = write_internal_template_file()
        uri = relative_uri(path)
        existing = CbamExportTemplate(
            id=uuid.UUID("d1000000-0000-4000-8000-000000000001"),
            organization_id=None,
            code=INTERNAL_TEMPLATE_CODE,
            name="EcoTrace SKDM Internal Development Template",
            template_type="INTERNAL_SKDM",
            version=INTERNAL_TEMPLATE_VERSION,
            mapping_version=INTERNAL_MAPPING_VERSION,
            storage_uri=uri,
            checksum=checksum,
            status="ACTIVE",
            description=(
                "INTERNAL DEVELOPMENT TEMPLATE. NOT AN OFFICIAL CBAM SUBMISSION FORMAT. "
                "Official CBAM workbook mapping is BLOCKED pending domain-expert delivery."
            ),
            activated_at=datetime.now(UTC),
        )
        db.add(existing)
        db.flush()

    from openpyxl import load_workbook

    path = resolve_uri(existing.storage_uri)
    wb = load_workbook(path)
    missing = [name for name in SHEETS if name not in wb.sheetnames]
    if missing:
        raise ValidationAppError(f"Internal template missing sheets: {missing}")

    existing_codes = {
        m.mapping_code
        for m in db.execute(
            select(CbamExportMapping).where(CbamExportMapping.export_template_id == existing.id)
        ).scalars()
    }
    for spec in internal_mapping_specs():
        code = str(spec["mapping_code"])
        if code in existing_codes:
            continue
        db.add(
            CbamExportMapping(
                export_template_id=existing.id,
                mapping_code=code,
                source_type=str(spec["source_type"]),
                source_path=str(spec["source_path"]),
                worksheet_name=str(spec["worksheet_name"]),
                destination_type=str(spec["destination_type"]),
                destination_reference=str(spec["destination_reference"]),
                value_type=str(spec["value_type"]),
                required=bool(spec["required"]),
                transformation_code=str(spec["transformation_code"])
                if spec.get("transformation_code")
                else None,
                notes=str(spec.get("notes") or ""),
            )
        )
    db.flush()
    return existing


def list_export_templates(
    db: Session,
    user: User,
    organization_id: uuid.UUID,
    *,
    page: int,
    page_size: int,
) -> Page[ExportTemplateResponse]:
    require_cbam_view(db, user, organization_id)
    ensure_internal_export_template(db)
    stmt = select(CbamExportTemplate).where(
        or_(
            CbamExportTemplate.organization_id.is_(None),
            CbamExportTemplate.organization_id == organization_id,
        ),
        CbamExportTemplate.status != "ARCHIVED",
    )
    total = db.execute(select(func.count()).select_from(stmt.subquery())).scalar_one()
    rows = list(
        db.execute(
            stmt.order_by(CbamExportTemplate.code.asc(), CbamExportTemplate.version.desc())
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
        .scalars()
        .all()
    )
    return paginate(
        [_to_template(r) for r in rows],
        page=page,
        page_size=page_size,
        total_items=int(total),
    )


def get_export_template(
    db: Session,
    user: User,
    organization_id: uuid.UUID,
    template_id: uuid.UUID,
) -> ExportTemplateResponse:
    require_cbam_view(db, user, organization_id)
    ensure_internal_export_template(db)
    row = db.execute(
        select(CbamExportTemplate).where(
            CbamExportTemplate.id == template_id,
            or_(
                CbamExportTemplate.organization_id.is_(None),
                CbamExportTemplate.organization_id == organization_id,
            ),
        )
    ).scalar_one_or_none()
    if row is None:
        raise NotFoundError("Export template not found.")
    return _to_template(row)


def get_active_template_for_org(
    db: Session, organization_id: uuid.UUID, template_id: uuid.UUID | None = None
) -> CbamExportTemplate:
    ensure_internal_export_template(db)
    if template_id is not None:
        row = db.execute(
            select(CbamExportTemplate).where(
                CbamExportTemplate.id == template_id,
                CbamExportTemplate.status == "ACTIVE",
                or_(
                    CbamExportTemplate.organization_id.is_(None),
                    CbamExportTemplate.organization_id == organization_id,
                ),
            )
        ).scalar_one_or_none()
        if row is None:
            raise NotFoundError("Active export template not found.")
        return row
    row = db.execute(
        select(CbamExportTemplate)
        .where(
            CbamExportTemplate.status == "ACTIVE",
            CbamExportTemplate.template_type == "INTERNAL_SKDM",
            CbamExportTemplate.organization_id.is_(None),
        )
        .order_by(CbamExportTemplate.activated_at.desc().nulls_last())
        .limit(1)
    ).scalar_one_or_none()
    if row is None:
        raise ValidationAppError("No active internal export template is available.")
    return row


def list_template_mappings(db: Session, template_id: uuid.UUID) -> list[CbamExportMapping]:
    return list(
        db.execute(
            select(CbamExportMapping)
            .where(CbamExportMapping.export_template_id == template_id)
            .order_by(CbamExportMapping.mapping_code.asc())
        ).scalars()
    )


def verify_template_file(template: CbamExportTemplate) -> Path:
    path = resolve_uri(template.storage_uri)
    if not path.is_file():
        raise ValidationAppError("Export template file is missing.")
    if path.suffix.lower() != ".xlsx":
        raise ValidationAppError("Only .xlsx templates are supported (.xlsm rejected).")
    checksum = sha256_file(path)
    if checksum != template.checksum:
        raise ValidationAppError("Export template checksum mismatch.")
    return path
