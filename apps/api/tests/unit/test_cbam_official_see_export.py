"""Phase 12A Official SEE export unit tests (LibreOffice-free where possible)."""

from __future__ import annotations

import shutil
import tempfile
import uuid
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import load_workbook

from ecotrace.modules.cbam.application.official_see_export.capacity import (
    CapacityUsage,
    assess_capacity,
)
from ecotrace.modules.cbam.application.official_see_export.clearing import (
    clear_example_and_unused_inputs,
)
from ecotrace.modules.cbam.application.official_see_export.constants import (
    CAPACITY_FUEL_ACTIVITIES,
    CAPACITY_PROCESSES,
    CODE_CAPACITY_EXCEEDED_PROCESSES,
    CODE_RECALCULATION_ENGINE_UNAVAILABLE,
    EXAMPLE_IDENTIFIERS,
    MAPPING_VERSION,
    TEMPLATE_SHA256,
)
from ecotrace.modules.cbam.application.official_see_export.context import (
    OfficialSeeExportContext,
)
from ecotrace.modules.cbam.application.official_see_export.forensics import (
    collect_formula_map,
    inspect_workbook,
)
from ecotrace.modules.cbam.application.official_see_export.leakage import scan_example_leakage
from ecotrace.modules.cbam.application.official_see_export.mapping.loader import load_manifest
from ecotrace.modules.cbam.application.official_see_export.recalc import (
    RecalculationEngineUnavailable,
    soffice_available,
)
from ecotrace.modules.cbam.application.official_see_export.security import (
    excel_safe_text,
    safe_export_filename,
)
from ecotrace.modules.cbam.application.official_see_export.writer import (
    write_official_see_workbook,
)

TEMPLATE = Path(__file__).resolve().parents[4] / (
    "local-reference/CBAM SEE V2.1_Example Steel 3 Screws and nuts_final "
    "Dosyasının Kopyası- (1) (1).xlsx"
)


pytestmark = pytest.mark.skipif(
    not TEMPLATE.is_file(), reason="Official SEE template not present in local-reference"
)


def _minimal_ctx(**overrides: object) -> OfficialSeeExportContext:
    base = OfficialSeeExportContext(
        organization_id=uuid.uuid4(),
        organization_name="EcoTrace Test Org",
        binding_id=uuid.uuid4(),
        installation={"name": "Test Installation"},
        reporting_period={
            "start_date": "2026-01-01",
            "end_date": "2026-03-31",
            "label": "2026Q1",
        },
        goods=[{"goods_type": "Iron or steel products"}],
        processes=[
            {
                "slot_index": 0,
                "id": str(uuid.uuid4()),
                "name": "EcoTrace Process A",
                "goods_type": "Iron or steel products",
                "produced_quantity": 10.0,
                "marketed_quantity": 10.0,
                "non_cbam_quantity": 0.0,
                "allocated_direct_tco2e": 1.5,
                "allocated_electricity_mwh": 2.0,
                "electricity_ef": 0.4,
                "electricity_ef_source": "D.4(b)",
                "exported_electricity_mwh": None,
                "exported_electricity_ef": None,
                "has_measurable_heat": False,
                "has_waste_gas": False,
                "heat_import_tj": None,
                "heat_export_tj": None,
                "heat_import_ef": None,
                "heat_export_ef": None,
                "waste_import_tj": None,
                "waste_export_tj": None,
                "product_uses": [],
                "product_use_count": 0,
            }
        ],
        precursors=[],
        fuels=[],
        pee_products=[
            {
                "slot_index": 0,
                "name": "EcoTrace Process A",
                "cn_code": "73181595",
                "description": None,
                "specific_direct": 0.15,
                "specific_indirect": 0.08,
                "specific_total": 0.23,
            }
        ],
        pee_result_id=uuid.uuid4(),
        dea_result_id=uuid.uuid4(),
        iea_result_id=uuid.uuid4(),
        expected_outputs={
            ("Summary_Products", "I10"): 0.15,
            ("Summary_Products", "J10"): 0.08,
            ("Summary_Products", "K10"): 0.23,
        },
        source_fingerprint="abc",
    )
    for key, value in overrides.items():
        setattr(base, key, value)
    return base


def test_manifest_loads_and_matches_template_hash() -> None:
    manifest = load_manifest()
    assert manifest.mapping_version == MAPPING_VERSION
    assert manifest.template_sha256 == TEMPLATE_SHA256
    assert len(manifest.entries) > 100
    assert any(e.direction == "CLEAR_EXAMPLE" for e in manifest.entries)
    assert any(e.direction == "FORMULA" for e in manifest.entries)
    assert any(e.direction == "OUTPUT" for e in manifest.entries)
    assert manifest.capacities.get("fuel_activities") == CAPACITY_FUEL_ACTIVITIES


def test_capacity_exceeded_typed_codes() -> None:
    codes = assess_capacity(
        CapacityUsage(
            installations=1,
            goods=1,
            processes=CAPACITY_PROCESSES + 1,
            precursors=0,
            max_process_product_uses=0,
            max_precursor_product_uses=0,
            fuel_activities=0,
        )
    )
    assert CODE_CAPACITY_EXCEEDED_PROCESSES in codes


def test_clearing_removes_urun_ham_madde_firma() -> None:
    manifest = load_manifest()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "work.xlsx"
        shutil.copy2(TEMPLATE, out)
        wb = load_workbook(out)
        assert wb["A_InstData"]["L83"].value == "Ürün 1"
        assert wb["A_InstData"]["L102"].value == "Ham madde 1"
        assert wb["A_InstData"]["I20"].value == "Firma"
        clear_example_and_unused_inputs(wb, manifest, used_input_keys=set())
        assert wb["A_InstData"]["L83"].value is None
        assert wb["A_InstData"]["L102"].value is None
        assert wb["A_InstData"]["I20"].value is None
        leaks = scan_example_leakage(wb, manifest, used_input_keys=set())
        example_leaks = [f for f in leaks if f.kind == "EXAMPLE_IDENTIFIER"]
        assert example_leaks == []


def test_writer_preserves_formulas_and_clears_examples() -> None:
    manifest = load_manifest()
    before = load_workbook(TEMPLATE)
    before_formulas = collect_formula_map(before)
    ctx = _minimal_ctx()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "see.xlsx"
        meta = write_official_see_workbook(
            template_path=TEMPLATE, output_path=out, ctx=ctx, manifest=manifest
        )
        assert meta["formulaCount"] == len(before_formulas)
        after = load_workbook(out)
        after_formulas = collect_formula_map(after)
        assert len(after_formulas) == len(before_formulas)
        assert after["A_InstData"]["I20"].value == "Test Installation"
        assert after["A_InstData"]["L83"].value == "EcoTrace Process A"
        assert after["D_Processes"]["L24"].value is not None
        assert str(after["D_Processes"]["L24"].value).startswith("=")
        # Example identifiers gone from master slots
        for needle in ("Ürün 1", "Ham madde 1", "Firma"):
            assert needle not in EXAMPLE_IDENTIFIERS or after["A_InstData"]["L83"].value != needle


def test_data_validations_and_named_ranges_preserved() -> None:
    forensics_before = inspect_workbook(TEMPLATE)
    ctx = _minimal_ctx()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "see.xlsx"
        write_official_see_workbook(template_path=TEMPLATE, output_path=out, ctx=ctx)
        forensics_after = inspect_workbook(out)
    assert forensics_after.named_range_count == forensics_before.named_range_count
    assert "InputOutput" in forensics_after.sheet_names
    hidden = {s.title: s.state for s in forensics_after.sheets}
    assert hidden.get("InputOutput") == "hidden"
    assert hidden.get("Translations") == "hidden"
    a_before = next(s for s in forensics_before.sheets if s.title == "A_InstData")
    a_after = next(s for s in forensics_after.sheets if s.title == "A_InstData")
    assert a_after.data_validation_count == a_before.data_validation_count


def test_package_writer_reload_preserves_formulas_without_openpyxl_save() -> None:
    """Read-only openpyxl load is OK; never round-trip save (drops CF extLst)."""
    ctx = _minimal_ctx()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "see.xlsx"
        write_official_see_workbook(template_path=TEMPLATE, output_path=out, ctx=ctx)
        wb = load_workbook(out, data_only=False)
        assert wb["A_InstData"]["I20"].value == "Test Installation"
        assert str(wb["D_Processes"]["L24"].value).startswith("=")


def test_formula_injection_protection() -> None:
    assert excel_safe_text("=CMD()") == "'=CMD()"
    assert excel_safe_text("+1+1") == "'+1+1"
    assert excel_safe_text("-1") == "'-1"
    assert excel_safe_text("@SUM") == "'@SUM"
    assert excel_safe_text("safe") == "safe"
    name = safe_export_filename(
        installation_name="Plant/A",
        period_label="2026 Q1",
        generated_on=date(2026, 8, 30),
    )
    assert name.startswith("CBAM_SEE_")
    assert "/" not in name
    assert name.endswith(".xlsx")


def test_tax_fields_preserved_not_overwritten_with_zeros() -> None:
    """F_Tools tax/benchmark formulas stay formulas; first-release does not invent tax inputs."""
    manifest = load_manifest()
    tax_preserve = [
        e
        for e in manifest.entries
        if e.sheet == "F_Tools" and e.direction in {"PRESERVE", "FORMULA"}
    ]
    assert tax_preserve
    ctx = _minimal_ctx()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "see.xlsx"
        write_official_see_workbook(
            template_path=TEMPLATE, output_path=out, ctx=ctx, manifest=manifest
        )
        wb = load_workbook(out)
        # Sample formula cells on F_Tools remain formulas
        for entry in tax_preserve[:5]:
            if entry.expected_formula:
                assert wb[entry.sheet][entry.cell].value == entry.expected_formula


def test_biogenic_left_blank_policy_in_manifest() -> None:
    manifest = load_manifest()
    bio = [e for e in manifest.entries if "biogenic" in e.semantic_field.lower()]
    assert bio
    assert all(e.direction in {"CLEAR_EXAMPLE", "PRESERVE", "INPUT"} for e in bio)


def test_recalc_engine_unavailable_path() -> None:
    with patch(
        "ecotrace.modules.cbam.application.official_see_export.recalc.resolve_soffice_path",
        return_value=None,
    ):
        assert soffice_available() is False
        from ecotrace.modules.cbam.application.official_see_export.recalc import (
            recalculate_workbook,
        )

        with pytest.raises(RecalculationEngineUnavailable) as exc:
            recalculate_workbook(TEMPLATE)
        assert exc.value.code == CODE_RECALCULATION_ENGINE_UNAVAILABLE


def test_readiness_blocks_when_stale_mocked() -> None:
    from ecotrace.modules.cbam.application.official_see_export import readiness as readiness_mod
    from ecotrace.modules.cbam.application.official_see_export.constants import (
        CODE_DEA_MISSING_OR_STALE,
        CODE_PEE_V2_MISSING_OR_STALE,
    )

    db = MagicMock()
    user = MagicMock()
    org_id = uuid.uuid4()
    binding_id = uuid.uuid4()

    with (
        patch.object(readiness_mod, "require_cbam_view"),
        patch.object(readiness_mod, "get_binding_for_org", return_value=MagicMock()),
        patch.object(readiness_mod, "validate_manifest_or_raise"),
        patch.object(
            readiness_mod,
            "list_product_profiles_for_binding",
            return_value=MagicMock(items=[]),
        ),
        patch.object(
            readiness_mod,
            "get_monthly_production_basis_summary",
            return_value=MagicMock(allocation_basis_ready=True),
        ),
        patch.object(
            readiness_mod,
            "get_direct_emissions_allocation_summary",
            return_value=MagicMock(current_result_id=None, current_is_stale=False),
        ),
        patch.object(
            readiness_mod,
            "get_indirect_emissions_allocation_summary",
            return_value=MagicMock(current_result_id=None, current_is_stale=False),
        ),
        patch.object(
            readiness_mod,
            "get_production_process_binding_summary",
            return_value=MagicMock(processes=[]),
        ),
        patch.object(
            readiness_mod,
            "get_purchased_precursor_binding_summary",
            return_value=MagicMock(precursors=[], unbalanced_count=0),
        ),
        patch.object(
            readiness_mod,
            "get_product_embedded_emissions_readiness",
            return_value=MagicMock(
                status="NOT_READY",
                current_result_id=None,
                current_is_stale=True,
                blocking_issue_codes=[],
                stale_reason_codes=[],
            ),
        ),
        patch.object(readiness_mod, "soffice_available", return_value=False),
    ):
        db.execute.return_value.scalars.return_value.all.return_value = []
        result = readiness_mod.assess_official_see_readiness(db, user, org_id, binding_id)
    assert result.ready is False
    assert CODE_DEA_MISSING_OR_STALE in result.blocking_issue_codes
    assert CODE_PEE_V2_MISSING_OR_STALE in result.blocking_issue_codes


def test_migration_0030_chained_from_0029() -> None:
    path = (
        Path(__file__).resolve().parents[2]
        / "src"
        / "ecotrace"
        / "db"
        / "migrations"
        / "versions"
        / "0030_cbam_official_see_export.py"
    )
    source = path.read_text(encoding="utf-8")
    assert 'revision: str = "0030_cbam_official_see"' in source
    assert len("0030_cbam_official_see") <= 32
    assert 'down_revision: str | None = "0029_cbam_pee_v2"' in source
    assert "cbam_official_see_export_runs" in source
    assert "cbam_official_see_export_artifacts" in source
    assert "client_request_id" in source
    assert "source_fingerprint" in source
    assert "DROP TABLE IF EXISTS cbam_official_see_export_artifacts" in source


def test_method_selectors_never_process_emissions_or_mass_balance() -> None:
    from ecotrace.modules.cbam.application.official_see_export.constants import (
        UNSUPPORTED_METHODS,
        UNSUPPORTED_MONITORING_APPROACHES,
    )

    assert "PROCESS_EMISSIONS" in UNSUPPORTED_METHODS
    assert "MASS_BALANCE" in UNSUPPORTED_METHODS
    assert "Process emissions" in UNSUPPORTED_MONITORING_APPROACHES
    assert "Mass Balance" in UNSUPPORTED_MONITORING_APPROACHES
    manifest = load_manifest()
    assert any("combustion_only" in e.semantic_field for e in manifest.entries)
    # Tax calc / HYBRID / Process Emissions / Mass Balance are out of first-release scope.
    scope = " ".join(manifest.unsupported_scope).lower()
    assert "hybrid" in scope
    assert "process_emissions" in scope
    assert "mass_balance" in scope
    assert "tax_calculations" in scope


def test_package_preserves_cf_extlst_named_ranges_and_hidden_sheets() -> None:
    from ecotrace.modules.cbam.application.official_see_export.package_inventory import (
        assert_critical_markers_preserved,
        inventory_package,
    )

    before = inventory_package(TEMPLATE)
    ctx = _minimal_ctx()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "see.xlsx"
        write_official_see_workbook(template_path=TEMPLATE, output_path=out, ctx=ctx)
        after = inventory_package(out)
    assert_critical_markers_preserved(before, after)
    assert after.named_range_count == before.named_range_count
    states = {s.name: s.state for s in after.sheets}
    assert states.get("InputOutput") == "hidden"
    assert states.get("Translations") == "hidden"
    assert after.total_cf_rule == before.total_cf_rule
    assert after.total_ext_lst == before.total_ext_lst


def test_steel_mapping_write_clear_unused_slots_and_leakage() -> None:
    from ecotrace.modules.cbam.application.official_see_export.writer import (
        build_used_input_map,
    )

    manifest = load_manifest()
    steel_inputs = [
        e
        for e in manifest.by_direction("INPUT")
        if e.sheet == "Summary_Products"
        and e.semantic_field.endswith(
            (
                ".reducing_agent",
                ".steel_mill_identification_number",
                ".percent_mn",
                ".percent_cr",
                ".percent_ni",
                ".percent_other_alloys",
                ".percent_other_materials",
            )
        )
    ]
    assert len(steel_inputs) == 70  # 7 fields × 10 slots

    pee = {
        "slot_index": 0,
        "name": "EcoTrace Process A",
        "cn_code": "73082000",
        "description": None,
        "specific_direct": 0.15,
        "specific_indirect": 0.08,
        "specific_total": 0.23,
        "reducing_agent": "Coal or coke",
        "steel_mill_identification_number": "MILL-42",
        "percent_mn": 0.0,  # explicit zero
        "percent_cr": 1.25,
        "percent_ni": None,  # blank
        "percent_other_alloys": 0.5,
        "percent_other_materials": 0.0,
    }
    ctx = _minimal_ctx(pee_products=[pee])
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "see.xlsx"
        write_official_see_workbook(
            template_path=TEMPLATE, output_path=out, ctx=ctx, manifest=manifest
        )
        wb = load_workbook(out, data_only=False)
        sp = wb["Summary_Products"]
        assert sp["P10"].value == "Coal or coke"
        assert sp["Q10"].value == "MILL-42"
        assert sp["R10"].value == 0
        assert sp["S10"].value == 1.25
        assert sp["T10"].value is None
        assert sp["U10"].value == 0.5
        assert sp["X10"].value == 0
        # Unused slots cleared (example Natural gas gone)
        assert sp["P11"].value is None
        assert sp["Q11"].value is None
        used = set(build_used_input_map(ctx, manifest).keys())
        leaks = scan_example_leakage(wb, manifest, used_input_keys=used)
        assert leaks == []


def test_resolve_soffice_prefers_explicit_macos_path() -> None:
    from ecotrace.modules.cbam.application.official_see_export.recalc import (
        resolve_soffice_path,
    )

    explicit = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    resolved = resolve_soffice_path()
    if explicit.is_file():
        assert resolved == explicit


def test_libreoffice_recalc_and_parity_when_soffice_present() -> None:
    from ecotrace.modules.cbam.application.official_see_export.parity import (
        scan_output_formula_errors,
        workbook_places_for_cell,
    )
    from ecotrace.modules.cbam.application.official_see_export.recalc import (
        cleanup_recalc_dir,
        recalculate_workbook,
        resolve_soffice_path,
    )

    soffice = resolve_soffice_path()
    # NO skipif: if the known path exists this must run; otherwise fail closed loudly.
    assert soffice is not None, (
        "LibreOffice soffice required for Phase 12A+ acceptance "
        "(expected /Applications/LibreOffice.app/Contents/MacOS/soffice)"
    )
    assert workbook_places_for_cell("0.000") == 3

    pee = {
        "slot_index": 0,
        "name": "EcoTrace Process A",
        "cn_code": "73181595",
        "description": None,
        "specific_direct": 0.15,
        "specific_indirect": 0.08,
        "specific_total": 0.23,
        "reducing_agent": "Hydrogen",
        "steel_mill_identification_number": "M-LO",
        "percent_mn": 0.0,
        "percent_cr": None,
        "percent_ni": None,
        "percent_other_alloys": None,
        "percent_other_materials": None,
    }
    ctx = _minimal_ctx(pee_products=[pee])
    manifest = load_manifest()
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "see.xlsx"
        write_official_see_workbook(
            template_path=TEMPLATE, output_path=out, ctx=ctx, manifest=manifest
        )
        recalc = recalculate_workbook(out, timeout_seconds=300)
        try:
            errs = scan_output_formula_errors(recalc, manifest, product_count=1)
            # G10/G26 must not be #N/A when CN is in Parameters_CNCodes.
            assert errs == [], f"Unexpected formula errors: {errs[:5]}"
            wb = load_workbook(recalc, data_only=True)
            assert wb["Summary_Products"]["F10"].value == "73181595"
            assert wb["Summary_Products"]["P10"].value == "Hydrogen"
            assert wb["Summary_Products"]["R10"].value == 0
            g10 = wb["Summary_Products"]["G10"].value
            assert g10 is not None and not str(g10).startswith("#")
        finally:
            cleanup_recalc_dir(recalc)


def test_formula_error_tokens_fail_parity() -> None:
    from ecotrace.core.exceptions import BusinessRuleError
    from ecotrace.modules.cbam.application.official_see_export.constants import (
        CODE_FORMULA_PARITY_FAILED,
    )
    from ecotrace.modules.cbam.application.official_see_export.parity import (
        assert_parity_or_raise,
    )

    manifest = load_manifest()
    with tempfile.TemporaryDirectory() as tmp:
        # Craft a tiny workbook with an OUTPUT cell showing #N/A
        from openpyxl import Workbook

        path = Path(tmp) / "broken.xlsx"
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Summary_Products"
        ws["I10"] = "#N/A"
        wb.save(path)
        with pytest.raises(BusinessRuleError) as exc:
            assert_parity_or_raise(
                path,
                manifest,
                {("Summary_Products", "I10"): 0.15},
            )
        assert exc.value.code == CODE_FORMULA_PARITY_FAILED
        assert any(d.get("diff") == "formula_error" for d in (exc.value.details or []))


def test_workbook_places_for_cell_zero_dot_three() -> None:
    from ecotrace.modules.cbam.application.official_see_export.parity import (
        SEE_IJK_DECIMAL_PLACES,
        SEE_IJK_NUMBER_FORMAT,
        workbook_places_for_cell,
    )

    assert SEE_IJK_NUMBER_FORMAT == "0.000"
    assert SEE_IJK_DECIMAL_PLACES == 3
    assert workbook_places_for_cell("0.000") == 3
    assert workbook_places_for_cell("#,##0.000") == 3
    assert workbook_places_for_cell("0.00") == 2
    assert workbook_places_for_cell("General") == 0


def test_scan_detects_summary_communication_na() -> None:
    """Expanded scanner must catch Communication #N/A missed by OUTPUT-only scan."""
    from openpyxl import Workbook

    from ecotrace.modules.cbam.application.official_see_export.parity import (
        scan_output_formula_errors,
    )

    manifest = load_manifest()
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "comm_na.xlsx"
        wb = Workbook()
        sp = wb.active
        assert sp is not None
        sp.title = "Summary_Products"
        sp["D10"] = "EcoTrace Process A"
        sp["F10"] = "73181595"
        sp["I10"] = 0.1
        sp["J10"] = 0.2
        sp["K10"] = 0.3
        sp["P10"] = "Natural gas"
        sc = wb.create_sheet("Summary_Communication")
        sc["F26"] = "73181595"
        sc["G26"] = "#N/A"
        sc["I26"] = 0.1
        sc["J26"] = 0.2
        sc["K26"] = 0.3
        wb.save(path)

        errs = scan_output_formula_errors(path, manifest, product_count=1)
        assert errs, "Expected Communication #N/A to be detected"
        assert any(
            e.sheet == "Summary_Communication" and e.cell == "G26" and e.diff == "formula_error"
            for e in errs
        )


def test_lo_package_allowlist_accepts_calcchain_drop() -> None:
    from ecotrace.modules.cbam.application.official_see_export.package_inventory import (
        LO_REGENERATED_OR_OPTIONAL_PARTS,
        PackageInventory,
        SheetPackageStats,
        assert_package_acceptable_after_libreoffice,
        part_is_lo_allowlisted,
    )

    assert "xl/calcChain.xml" in LO_REGENERATED_OR_OPTIONAL_PARTS
    assert part_is_lo_allowlisted("xl/calcChain.xml")
    assert part_is_lo_allowlisted("xl/printerSettings/printerSettings1.bin")
    assert part_is_lo_allowlisted("xl/worksheets/_rels/sheet13.xml.rels")
    assert not part_is_lo_allowlisted("xl/worksheets/sheet1.xml")

    sheet = SheetPackageStats(
        name="Summary_Products",
        path="xl/worksheets/sheet1.xml",
        state="visible",
        cf_rule_count=2,
        ext_lst_count=1,
        data_validation_count=1,
        formula_count=10,
        has_drawing_rel=False,
    )
    before = PackageInventory(
        part_names=(
            "xl/workbook.xml",
            "xl/styles.xml",
            "xl/calcChain.xml",
            "xl/worksheets/sheet1.xml",
        ),
        part_sha256={
            "xl/workbook.xml": "a",
            "xl/styles.xml": "b",
            "xl/calcChain.xml": "c",
            "xl/worksheets/sheet1.xml": "d",
        },
        sheets=(sheet,),
        named_range_count=5,
        named_ranges=("N1",),
        total_cf_rule=2,
        total_ext_lst=1,
        drawing_parts=(),
    )
    after = PackageInventory(
        part_names=("xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet1.xml"),
        part_sha256={
            "xl/workbook.xml": "a2",
            "xl/styles.xml": "b",
            "xl/worksheets/sheet1.xml": "d2",
        },
        sheets=(sheet,),
        named_range_count=5,
        named_ranges=("N1",),
        total_cf_rule=2,
        total_ext_lst=1,
        drawing_parts=(),
    )
    assert_package_acceptable_after_libreoffice(before, after)

    broken = PackageInventory(
        part_names=("xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet1.xml"),
        part_sha256={
            "xl/workbook.xml": "a2",
            "xl/styles.xml": "b",
            "xl/worksheets/sheet1.xml": "d2",
        },
        sheets=(
            SheetPackageStats(
                name="Summary_Products",
                path="xl/worksheets/sheet1.xml",
                state="visible",
                cf_rule_count=0,
                ext_lst_count=1,
                data_validation_count=1,
                formula_count=10,
                has_drawing_rel=False,
            ),
        ),
        named_range_count=5,
        named_ranges=("N1",),
        total_cf_rule=0,
        total_ext_lst=1,
        drawing_parts=(),
    )
    with pytest.raises(AssertionError, match="cfRule"):
        assert_package_acceptable_after_libreoffice(before, broken)
