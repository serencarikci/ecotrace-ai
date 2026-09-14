"""Phase 6A CN catalog extraction, seed, and resolution tests."""

from __future__ import annotations

import os
import uuid
from pathlib import Path

import pytest
from sqlalchemy import func, select

from ecotrace.core.exceptions import BusinessRuleError, ConflictError, NotFoundError
from ecotrace.db.seed import DEMO_ORG_SLUG
from ecotrace.modules.cbam.application.cn_catalog_seed import (
    compute_seed_content_checksum,
    ensure_platform_cn_catalog,
    load_cn_catalog_seed_payload,
)
from ecotrace.modules.cbam.application.cn_catalog_service import (
    get_cn_code,
    list_cn_codes,
    list_controlled_list_values,
    normalize_cn_code,
    resolve_cn_code,
)
from ecotrace.modules.cbam.infrastructure.models import CbamCnCode, CbamCnCodeDataset
from ecotrace.modules.identity.infrastructure.models import User
from ecotrace.modules.organizations.infrastructure.models import Organization

_REPO_ROOT = Path(__file__).resolve().parents[4]
_DEFAULT_SEE_WORKBOOK = (
    _REPO_ROOT
    / 'local-reference'
    / (
        'CBAM SEE V2.1_Example Steel 3 Screws and nuts_final '
        'Dosyasının Kopyası- (1) (1).xlsx'
    )
)
WORKBOOK = Path(os.environ.get('CBAM_SEE_WORKBOOK_PATH', str(_DEFAULT_SEE_WORKBOOK)))


def _org(db):
    return db.execute(select(Organization).where(Organization.slug == DEMO_ORG_SLUG)).scalar_one()


def _admin(db):
    return db.execute(
        select(User).where(User.normalized_email == 'orgadmin@ecotrace.dev')
    ).scalar_one()


def _viewer(db):
    return db.execute(
        select(User).where(User.normalized_email == 'viewer@ecotrace.dev')
    ).scalar_one()


@pytest.mark.skipif(not WORKBOOK.is_file(), reason='CBAM SEE workbook not available')
def test_source_workbook_readable_and_sheets_present() -> None:
    from openpyxl import load_workbook

    assert WORKBOOK.is_file()
    assert WORKBOOK.stat().st_size > 0
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    for sheet in (
        'Parameters_CNCodes',
        'c_CodeLists',
        'Summary_Products',
        'D_Processes',
    ):
        assert sheet in wb.sheetnames
    wb.close()


def test_seed_payload_preserves_text_and_leading_structure() -> None:
    payload = load_cn_catalog_seed_payload()
    codes = payload['cnCodes']
    assert len(codes) == 569
    assert all(isinstance(c['normalizedCode'], str) for c in codes)
    # Representative cement / clay codes keep leading structure as text.
    clay = next(c for c in codes if c['normalizedCode'] == '25070080')
    assert clay['displayCode'] == '2507 00 80'
    assert clay['cnKey'].endswith('0080')
    steel = [c for c in codes if c['cbamSector'] == 'Iron or steel products']
    assert len(steel) > 100
    screws = next(c for c in codes if c['normalizedCode'] == '73181595')
    assert 'Screws and bolts' in screws['descriptionEn']
    nuts = next(c for c in codes if c['normalizedCode'] == '73181699')
    assert 'Nuts' in nuts['descriptionEn']
    assert payload['controlledLists']['REDUCING_AGENT'] == [
        {'code': 'Coal or coke', 'label': 'Coal or coke'},
        {'code': 'Natural gas', 'label': 'Natural gas'},
        {'code': 'Biogas', 'label': 'Biogas'},
        {'code': 'Hydrogen', 'label': 'Hydrogen'},
    ]
    # No duplicate normalized codes in this workbook version.
    norms = [c['normalizedCode'] for c in codes]
    assert len(norms) == len(set(norms))
    checksum = compute_seed_content_checksum(payload)
    assert checksum == payload['dataset']['contentChecksum']
    assert checksum == compute_seed_content_checksum(payload)
    # Phase 6A+: every CN carries explicit fieldApplicability (all keys, booleans).
    expected_keys = {
        'reducingAgent',
        'steelMillIdentificationNumber',
        'percentMn',
        'percentCr',
        'percentNi',
        'percentOtherAlloys',
        'percentOtherMaterials',
    }
    for code in codes:
        fa = code['fieldApplicability']
        assert set(fa) == expected_keys
        assert all(isinstance(fa[k], bool) for k in expected_keys)
    assert all(screws['fieldApplicability'][k] is True for k in expected_keys)
    cement = next(c for c in codes if c['normalizedCode'] == '25232900')
    assert all(cement['fieldApplicability'][k] is False for k in expected_keys)


def test_seed_idempotent_and_immutable_conflict(seeded_db) -> None:
    first = ensure_platform_cn_catalog(seeded_db)
    count1 = seeded_db.execute(select(func.count()).select_from(CbamCnCode)).scalar_one()
    second = ensure_platform_cn_catalog(seeded_db)
    count2 = seeded_db.execute(select(func.count()).select_from(CbamCnCode)).scalar_one()
    assert first.id == second.id
    assert count1 == count2 == 569

    original = first.content_checksum
    first.content_checksum = '0' * 64
    seeded_db.flush()
    with pytest.raises(ConflictError) as exc:
        ensure_platform_cn_catalog(seeded_db)
    assert exc.value.details[0]['code'] == 'IMMUTABLE_CN_DATASET_CONFLICT'
    first.content_checksum = original
    seeded_db.flush()


def test_new_dataset_version_can_coexist(seeded_db) -> None:
    ensure_platform_cn_catalog(seeded_db)
    other = CbamCnCodeDataset(
        id=uuid.uuid4(),
        dataset_code='CBAM_SEE_CN_CODES',
        dataset_version='SEE_V2.1_PATCH',
        content_checksum='a' * 64,
        source_workbook_name='test.xlsx',
        source_workbook_sha256='b' * 64,
        source_template_version='2.1-patch',
        valid_from=first_valid_from(seeded_db),
        valid_until=None,
        status='ACTIVE',
    )
    # Mark previous as superseded to allow multiple versions historically.
    current = seeded_db.execute(
        select(CbamCnCodeDataset).where(CbamCnCodeDataset.dataset_version == 'SEE_V2.1')
    ).scalar_one()
    current.status = 'SUPERSEDED'
    seeded_db.add(other)
    seeded_db.flush()
    versions = seeded_db.execute(
        select(CbamCnCodeDataset.dataset_version).order_by(CbamCnCodeDataset.dataset_version)
    ).scalars().all()
    assert 'SEE_V2.1' in versions and 'SEE_V2.1_PATCH' in versions


def first_valid_from(db):
    return db.execute(
        select(CbamCnCodeDataset.valid_from).where(CbamCnCodeDataset.dataset_version == 'SEE_V2.1')
    ).scalar_one()


def test_cn_resolve_search_filter(seeded_db) -> None:
    org = _org(seeded_db)
    admin = _admin(seeded_db)
    ensure_platform_cn_catalog(seeded_db)
    by_norm = resolve_cn_code(seeded_db, code='73181595')
    assert by_norm.normalized_code == '73181595'
    by_display = resolve_cn_code(seeded_db, code='7318 15 95')
    assert by_display.id == by_norm.id
    page = list_cn_codes(seeded_db, admin, org.id, page=1, page_size=20, q='73181595')
    assert page.total_items >= 1
    desc = list_cn_codes(seeded_db, admin, org.id, page=1, page_size=20, q='Screws and bolts')
    assert desc.total_items >= 1
    sector = list_cn_codes(
        seeded_db, admin, org.id, page=1, page_size=5, sector='Iron or steel products'
    )
    assert all(i.cbam_sector == 'Iron or steel products' for i in sector.items)
    with pytest.raises(NotFoundError):
        resolve_cn_code(seeded_db, code='00000000')
    detail = get_cn_code(seeded_db, admin, org.id, by_norm.id)
    assert detail.dataset_version == 'SEE_V2.1'
    assert detail.content_checksum
    agents = list_controlled_list_values(seeded_db, list_code='REDUCING_AGENT')
    assert [a.value_code for a in agents] == [
        'Coal or coke',
        'Natural gas',
        'Biogas',
        'Hydrogen',
    ]


def test_inactive_code_excluded_from_active_search(seeded_db) -> None:
    org = _org(seeded_db)
    admin = _admin(seeded_db)
    ensure_platform_cn_catalog(seeded_db)
    code = resolve_cn_code(seeded_db, code='73181595')
    code.status = 'INACTIVE'
    seeded_db.flush()
    page = list_cn_codes(seeded_db, admin, org.id, page=1, page_size=50, q='73181595')
    assert all(i.normalized_code != '73181595' for i in page.items)
    inactive = resolve_cn_code(seeded_db, code='73181595', allow_inactive=True)
    assert inactive.status == 'INACTIVE'


def test_ambiguous_cn_fails_closed(seeded_db) -> None:
    dataset = seeded_db.execute(
        select(CbamCnCodeDataset).where(CbamCnCodeDataset.dataset_version == 'SEE_V2.1')
    ).scalar_one()
    a = resolve_cn_code(seeded_db, code='25070080')
    b = CbamCnCode(
        id=uuid.uuid4(),
        dataset_id=dataset.id,
        cn_key='DUPLICATE_TEST_KEY',
        normalized_code='ZZZZZZZZ',
        display_code=a.display_code,
        description_en='dup',
        cbam_sector=a.cbam_sector,
        numbering_label=None,
        source_sheet='test',
        source_row=99999,
        status='ACTIVE',
        field_applicability=dict(a.field_applicability or {}),
    )
    seeded_db.add(b)
    seeded_db.flush()
    with pytest.raises(BusinessRuleError) as exc:
        resolve_cn_code(seeded_db, code=a.display_code)
    assert exc.value.details[0]['code'] == 'AMBIGUOUS_CN_CODE'


def test_normalize_cn_code_helper() -> None:
    assert normalize_cn_code('7318 15 95') == '73181595'


def test_viewer_can_search_cn_codes(seeded_db) -> None:
    org = _org(seeded_db)
    viewer = _viewer(seeded_db)
    page = list_cn_codes(seeded_db, viewer, org.id, page=1, page_size=10, q='7318')
    assert page.total_items >= 1


def test_cn_detail_exposes_authoritative_field_applicability(seeded_db) -> None:
    org = _org(seeded_db)
    admin = _admin(seeded_db)
    steel = resolve_cn_code(seeded_db, code='73181595')
    cement = resolve_cn_code(seeded_db, code='25232900')
    steel_detail = get_cn_code(seeded_db, admin, org.id, steel.id)
    cement_detail = get_cn_code(seeded_db, admin, org.id, cement.id)
    keys = {
        'reducingAgent',
        'steelMillIdentificationNumber',
        'percentMn',
        'percentCr',
        'percentNi',
        'percentOtherAlloys',
        'percentOtherMaterials',
    }
    steel_fa = steel_detail.field_applicability.model_dump(by_alias=True)
    cement_fa = cement_detail.field_applicability.model_dump(by_alias=True)
    assert set(steel_fa) == keys
    assert set(cement_fa) == keys
    assert all(isinstance(steel_fa[k], bool) for k in keys)
    assert all(steel_fa[k] is True for k in keys)
    assert all(cement_fa[k] is False for k in keys)
    # Persisted on the CN row — not reconstructed from sector name at read time.
    assert steel.field_applicability['reducingAgent'] is True
    assert cement.field_applicability['reducingAgent'] is False


def test_seeded_rows_persist_field_applicability(seeded_db) -> None:
    payload = load_cn_catalog_seed_payload()
    by_norm = {c['normalizedCode']: c['fieldApplicability'] for c in payload['cnCodes']}
    rows = list(seeded_db.execute(select(CbamCnCode)).scalars().all())
    assert len(rows) == 569
    for row in rows:
        assert row.field_applicability == by_norm[row.normalized_code]
