"""Phase 9A Conventional production-process tests."""

from __future__ import annotations

import os
import uuid
from decimal import Decimal
from pathlib import Path

import pytest
from sqlalchemy import text
from tests.cbam_dea_helpers import admin, org, setup_binding
from tests.cbam_profile_helpers import create_active_ready_profile, ensure_org_product

from ecotrace.core.exceptions import BusinessRuleError, NotFoundError, ValidationAppError
from ecotrace.modules.cbam.application import production_process_service
from ecotrace.modules.cbam.application.production_process_constants import (
    BALANCE_BALANCED,
    BALANCE_UNBALANCED,
    CODE_DIRECT_EMISSIONS_ALLOCATION_NOT_READY,
    CODE_HEAT_FIELDS_NOT_ALLOWED,
    CODE_INDIRECT_EMISSIONS_ALLOCATION_NOT_READY,
    CODE_PROCESS_METHOD_UNSUPPORTED,
    CODE_PRODUCT_DISTRIBUTION_UNBALANCED,
    CODE_WASTE_GAS_FIELDS_NOT_ALLOWED,
    CONST_EF_NAT_GAS_TCO2_PER_TJ,
    DATA_QUALITY_LIST_CODE,
    METHOD_CONVENTIONAL,
    METHOD_MASS_BALANCE,
    METHOD_PROCESS_EMISSIONS,
    READINESS_EMPTY,
    READINESS_UNBALANCED,
    WASTE_GAS_EXPORT_FACTOR,
    WORKBOOK_SHA256,
)
from ecotrace.modules.cbam.application.production_process_controlled_lists import (
    PROCESS_FIELD_MAP,
    data_quality_codes,
    list_production_process_controlled_lists,
)
from ecotrace.modules.cbam.application.production_process_math import (
    compute_distribution_balance,
    compute_exported_electricity_attribution,
    compute_measurable_heat_attribution,
    compute_waste_gas_attribution,
)
from ecotrace.modules.cbam.application.production_process_service import (
    ProductionProcessCreate,
    ProductionProcessUpdate,
    ProductionProcessVersionRequest,
    ProductUseCreate,
)
from ecotrace.modules.cbam.infrastructure.models import CbamProductProfileVersion
from ecotrace.modules.organizations.infrastructure.models import Organization

_REPO_ROOT = Path(__file__).resolve().parents[4]
_DEFAULT_SEE_WORKBOOK = (
    _REPO_ROOT
    / "local-reference"
    / ("CBAM SEE V2.1_Example Steel 3 Screws and nuts_final Dosyasının Kopyası- (1) (1).xlsx")
)
WORKBOOK = Path(os.environ.get("CBAM_SEE_WORKBOOK_PATH", str(_DEFAULT_SEE_WORKBOOK)))


@pytest.mark.skipif(not WORKBOOK.is_file(), reason="CBAM SEE workbook not available")
def test_workbook_d_processes_field_map_and_lists() -> None:
    import hashlib

    from openpyxl import load_workbook

    digest = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()
    assert digest == WORKBOOK_SHA256
    wb = load_workbook(WORKBOOK, data_only=False)
    assert "D_Processes" in wb.sheetnames
    ws = wb["D_Processes"]
    assert ws["L24"].value == '=IF(G11="","",SUM(L16:L23))'
    assert "SUM(L27,L32:L41)" in str(ws["L42"].value)
    assert "CONST_EFNatGas" in str(ws["T62"].value)
    assert "L57*L58" in str(ws["T58"].value).replace(" ", "")
    tr = wb["Translations"]
    assert tr.cell(1875, 3).value == "Mostly measurements & analyses"
    assert "produced_quantity_total" in PROCESS_FIELD_MAP
    codes = data_quality_codes()
    assert "MOSTLY_MEASUREMENTS_AND_ANALYSES" in codes
    assert any(
        lst.list_code == DATA_QUALITY_LIST_CODE
        for lst in list_production_process_controlled_lists()
    )
    wb.close()


def test_conventional_supported_and_disabled_methods_rejected(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    binding, installation = setup_binding(db, user, organization)
    profile = create_active_ready_profile(db, user, organization.id)

    ok = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(
            installation_profile_id=installation.id,
            name="P1",
            calculation_method=METHOD_CONVENTIONAL,
            product_profile_version_id=profile.id,
            produced_quantity=Decimal("10"),
            produced_quantity_unit="t",
            marketed_quantity=Decimal("10"),
            marketed_quantity_unit="t",
            non_cbam_quantity=Decimal("0"),
            non_cbam_quantity_unit="t",
            has_measurable_heat=False,
            has_waste_gas=False,
        ),
    )
    assert ok.calculation_method == METHOD_CONVENTIONAL

    for method in (METHOD_PROCESS_EMISSIONS, METHOD_MASS_BALANCE):
        with pytest.raises(BusinessRuleError) as exc:
            production_process_service.create_production_process(
                db,
                user,
                organization.id,
                binding.id,
                ProductionProcessCreate(
                    installation_profile_id=installation.id,
                    name="Bad",
                    calculation_method=method,
                ),
            )
        assert CODE_PROCESS_METHOD_UNSUPPORTED in str(exc.value.details)


def test_same_org_profile_and_cross_org_rejected(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    binding, installation = setup_binding(db, user, organization)
    profile = create_active_ready_profile(db, user, organization.id)
    process = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(
            installation_profile_id=installation.id,
            name="P1",
            product_profile_version_id=profile.id,
        ),
    )
    assert process.product_profile_version_id == profile.id

    other = Organization(name="Other Org", slug=f"other-{uuid.uuid4().hex[:8]}", is_active=True)
    db.add(other)
    db.flush()
    foreign_product = ensure_org_product(db, other.id)
    foreign = CbamProductProfileVersion(
        organization_id=other.id,
        product_id=foreign_product.id,
        version=1,
        status="active",
        classification_ready=True,
        product_name="Foreign",
        cn_normalized_code="73181595",
        cn_display_code="7318 15 95",
    )
    db.add(foreign)
    db.flush()
    with pytest.raises(NotFoundError):
        production_process_service.create_production_process(
            db,
            user,
            organization.id,
            binding.id,
            ProductionProcessCreate(
                installation_profile_id=installation.id,
                name="Cross",
                product_profile_version_id=foreign.id,
            ),
        )


def test_distribution_market_only_mixed_non_cbam_balance(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    binding, installation = setup_binding(db, user, organization)
    profile = create_active_ready_profile(db, user, organization.id)
    target = create_active_ready_profile(
        db, user, organization.id, product=ensure_org_product(db, organization.id, code="TGT")
    )

    p1 = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(
            installation_profile_id=installation.id,
            name="MarketOnly",
            product_profile_version_id=profile.id,
            produced_quantity=Decimal("39.34"),
            produced_quantity_unit="t",
            marketed_quantity=Decimal("39.34"),
            marketed_quantity_unit="t",
            non_cbam_quantity=Decimal("0"),
            non_cbam_quantity_unit="t",
            has_measurable_heat=False,
            has_waste_gas=False,
        ),
    )
    assert p1.distribution.balance_status == BALANCE_BALANCED
    assert p1.distribution.remaining_tonnes == Decimal("0")
    assert p1.distribution.all_to_market is True

    p2 = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(
            installation_profile_id=installation.id,
            name="Mixed",
            product_profile_version_id=profile.id,
            produced_quantity=Decimal("100"),
            produced_quantity_unit="t",
            marketed_quantity=Decimal("40"),
            marketed_quantity_unit="t",
            non_cbam_quantity=Decimal("10"),
            non_cbam_quantity_unit="t",
            has_measurable_heat=False,
            has_waste_gas=False,
        ),
    )
    production_process_service.create_product_use(
        db,
        user,
        organization.id,
        binding.id,
        p2.id,
        ProductUseCreate(
            target_product_profile_version_id=target.id,
            quantity=Decimal("50"),
            unit="t",
        ),
    )
    detail = production_process_service.get_production_process(
        db, user, organization.id, binding.id, p2.id
    )
    assert detail.distribution.balance_status == BALANCE_BALANCED
    assert detail.distribution.other_cbam_tonnes == Decimal("50")
    assert detail.distribution.non_cbam_tonnes == Decimal("10")


def test_unbalanced_draft_save_and_readiness_blocked(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    binding, installation = setup_binding(db, user, organization)
    profile = create_active_ready_profile(db, user, organization.id)

    process = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(
            installation_profile_id=installation.id,
            name="Unbalanced",
            product_profile_version_id=profile.id,
            produced_quantity=Decimal("100"),
            produced_quantity_unit="t",
            marketed_quantity=Decimal("60"),
            marketed_quantity_unit="t",
            non_cbam_quantity=Decimal("0"),
            non_cbam_quantity_unit="t",
            has_measurable_heat=False,
            has_waste_gas=False,
        ),
    )
    assert process.status == "draft"
    assert process.distribution.balance_status == BALANCE_UNBALANCED
    assert process.distribution.remaining_tonnes == Decimal("40")
    assert process.readiness.status == READINESS_UNBALANCED
    assert CODE_PRODUCT_DISTRIBUTION_UNBALANCED in process.readiness.blocking_issue_codes

    updated = production_process_service.update_production_process_draft(
        db,
        user,
        organization.id,
        binding.id,
        process.id,
        ProductionProcessUpdate(
            row_version=process.row_version,
            marketed_quantity=Decimal("70"),
        ),
    )
    assert updated.distribution.remaining_tonnes == Decimal("30")
    assert updated.readiness.status == READINESS_UNBALANCED


def test_decimal_unit_kg_to_tonnes_exact(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    binding, installation = setup_binding(db, user, organization)
    profile = create_active_ready_profile(db, user, organization.id)
    process = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(
            installation_profile_id=installation.id,
            name="Kg",
            product_profile_version_id=profile.id,
            produced_quantity=Decimal("1000"),
            produced_quantity_unit="kg",
            marketed_quantity=Decimal("1"),
            marketed_quantity_unit="t",
            non_cbam_quantity=Decimal("0"),
            non_cbam_quantity_unit="t",
            has_measurable_heat=False,
            has_waste_gas=False,
        ),
    )
    assert process.distribution.produced_tonnes == Decimal("1")
    assert process.distribution.balance_status == BALANCE_BALANCED


def test_allocation_missing_blocks_and_exported_electricity_readonly(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    binding, installation = setup_binding(db, user, organization)
    profile = create_active_ready_profile(db, user, organization.id)
    process = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(
            installation_profile_id=installation.id,
            name="Alloc",
            product_profile_version_id=profile.id,
            produced_quantity=Decimal("10"),
            produced_quantity_unit="t",
            marketed_quantity=Decimal("10"),
            marketed_quantity_unit="t",
            non_cbam_quantity=Decimal("0"),
            non_cbam_quantity_unit="t",
            has_measurable_heat=False,
            has_waste_gas=False,
        ),
    )
    assert CODE_DIRECT_EMISSIONS_ALLOCATION_NOT_READY in process.readiness.blocking_issue_codes
    assert CODE_INDIRECT_EMISSIONS_ALLOCATION_NOT_READY in process.readiness.blocking_issue_codes
    assert process.direct_emissions_allocation.product_allocated_value is None
    assert process.exported_electricity.source == "purchased_electricity_current"
    assert "Not subtracted" in process.exported_electricity.note


def test_heat_and_waste_conditional_validation(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    binding, installation = setup_binding(db, user, organization)

    with pytest.raises(ValidationAppError) as heat_exc:
        production_process_service.create_production_process(
            db,
            user,
            organization.id,
            binding.id,
            ProductionProcessCreate(
                installation_profile_id=installation.id,
                name="HeatBad",
                has_measurable_heat=False,
                heat_imported_quantity=Decimal("1"),
                heat_imported_unit="TJ",
            ),
        )
    assert CODE_HEAT_FIELDS_NOT_ALLOWED in str(heat_exc.value.details)

    with pytest.raises(ValidationAppError) as waste_exc:
        production_process_service.create_production_process(
            db,
            user,
            organization.id,
            binding.id,
            ProductionProcessCreate(
                installation_profile_id=installation.id,
                name="WasteBad",
                has_waste_gas=False,
                waste_gas_imported_quantity=Decimal("1"),
                waste_gas_imported_unit="TJ",
            ),
        )
    assert CODE_WASTE_GAS_FIELDS_NOT_ALLOWED in str(waste_exc.value.details)

    process = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(
            installation_profile_id=installation.id,
            name="HeatOk",
            has_measurable_heat=True,
            heat_imported_quantity=Decimal("2"),
            heat_imported_unit="TJ",
            heat_exported_quantity=Decimal("0.5"),
            heat_exported_unit="TJ",
            heat_imported_ef=Decimal("56.1"),
            heat_exported_ef=Decimal("56.1"),
            heat_ef_unit="tCO2/TJ",
            has_waste_gas=True,
            waste_gas_imported_quantity=Decimal("1"),
            waste_gas_imported_unit="TJ",
            waste_gas_exported_quantity=Decimal("0"),
            waste_gas_exported_unit="TJ",
        ),
    )
    assert process.measurable_heat.calculation_status == "CALCULATED"
    assert process.measurable_heat.attributed_tco2 == Decimal("2") * Decimal("56.1") - Decimal(
        "0.5"
    ) * Decimal("56.1")
    assert process.waste_gas.calculation_status == "CALCULATED"
    expected_wg = Decimal("1") * CONST_EF_NAT_GAS_TCO2_PER_TJ
    assert process.waste_gas.attributed_tco2 == expected_wg

    cleared = production_process_service.update_production_process_draft(
        db,
        user,
        organization.id,
        binding.id,
        process.id,
        ProductionProcessUpdate(row_version=process.row_version, has_measurable_heat=False),
    )
    assert cleared.measurable_heat.has_measurable_heat is False
    assert cleared.measurable_heat.imported_quantity is None


def test_controlled_lists_and_invalid_dq_code(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    meta = production_process_service.get_production_process_metadata(db, user, organization.id)
    dq = next(lst for lst in meta.controlled_lists if lst.list_code == DATA_QUALITY_LIST_CODE)
    assert {i.code for i in dq.items} == set(data_quality_codes())

    binding, installation = setup_binding(db, user, organization)
    with pytest.raises(ValidationAppError):
        production_process_service.create_production_process(
            db,
            user,
            organization.id,
            binding.id,
            ProductionProcessCreate(
                installation_profile_id=installation.id,
                name="DQ",
                data_quality_code="NOT_A_REAL_CODE",
            ),
        )


def test_empty_readiness_and_archive(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    binding, installation = setup_binding(db, user, organization)
    process = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(installation_profile_id=installation.id),
    )
    assert process.readiness.status == READINESS_EMPTY
    archived = production_process_service.archive_production_process(
        db,
        user,
        organization.id,
        binding.id,
        process.id,
        ProductionProcessVersionRequest(row_version=process.row_version),
    )
    assert archived.status == "archived"


def test_math_helpers_exact() -> None:
    bal = compute_distribution_balance(
        produced_tonnes=Decimal("10"),
        marketed_tonnes=Decimal("3"),
        other_cbam_tonnes=Decimal("4"),
        non_cbam_tonnes=Decimal("3"),
    )
    assert bal.balance_status == BALANCE_BALANCED
    under = compute_distribution_balance(
        produced_tonnes=Decimal("10"),
        marketed_tonnes=Decimal("3"),
        other_cbam_tonnes=Decimal("4"),
        non_cbam_tonnes=Decimal("2"),
    )
    assert under.remaining_tonnes == Decimal("1")
    assert under.balance_status == BALANCE_UNBALANCED
    heat = compute_measurable_heat_attribution(
        has_measurable_heat=True,
        imported_tj=Decimal("1"),
        exported_tj=Decimal("0"),
        imported_ef=Decimal("10"),
        exported_ef=Decimal("10"),
    )
    assert heat.attributed_tco2 == Decimal("10")
    waste = compute_waste_gas_attribution(
        has_waste_gas=True,
        imported_tj=Decimal("1"),
        exported_tj=Decimal("1"),
    )
    assert waste.attributed_tco2 == CONST_EF_NAT_GAS_TCO2_PER_TJ - (
        CONST_EF_NAT_GAS_TCO2_PER_TJ * WASTE_GAS_EXPORT_FACTOR
    )


def test_exported_electricity_t72_formula() -> None:
    calculated = compute_exported_electricity_attribution(
        has_exported_electricity=True,
        quantity_mwh=Decimal("2.5"),
        emission_factor=Decimal("0.4"),
    )
    assert calculated.status == "CALCULATED"
    assert calculated.attributed_direct_tco2e == Decimal("-1.0")
    assert calculated.formula_ref == "D_Processes!T72=-L71*L72"

    not_applicable = compute_exported_electricity_attribution(
        has_exported_electricity=False,
        quantity_mwh=None,
        emission_factor=None,
    )
    assert not_applicable.status == "NOT_APPLICABLE"
    assert not_applicable.attributed_direct_tco2e is None


def test_exported_electricity_incomplete_without_factor_or_provenance(seeded_db) -> None:
    db = seeded_db
    organization = org(db)
    user = admin(db)
    binding, installation = setup_binding(db, user, organization)
    product = ensure_org_product(db, organization.id, code=f"T72-{uuid.uuid4().hex[:6]}")
    profile = create_active_ready_profile(db, user, organization.id, product=product)
    process = production_process_service.create_production_process(
        db,
        user,
        organization.id,
        binding.id,
        ProductionProcessCreate(
            installation_profile_id=installation.id,
            name="Export process",
            product_profile_version_id=profile.id,
            produced_quantity=Decimal("10"),
            produced_quantity_unit="t",
            marketed_quantity=Decimal("10"),
            marketed_quantity_unit="t",
            non_cbam_quantity=Decimal("0"),
            non_cbam_quantity_unit="t",
            has_measurable_heat=False,
            has_waste_gas=False,
            has_exported_electricity=True,
            exported_electricity_quantity=Decimal("1"),
            exported_electricity_unit="MWh",
        ),
    )
    assert "EXPORTED_ELECTRICITY_FACTOR_REQUIRED" in process.readiness.blocking_issue_codes
    assert "EXPORTED_ELECTRICITY_PROVENANCE_REQUIRED" in process.readiness.blocking_issue_codes
    assert process.process_exported_electricity.calculation_status == "INCOMPLETE"
    assert process.process_exported_electricity.attributed_direct_tco2e is None


def test_production_process_migration_round_trip(seeded_db) -> None:
    db = seeded_db

    def downgrade() -> None:
        db.execute(text("DROP TABLE IF EXISTS cbam_production_process_product_uses CASCADE"))
        db.execute(text("DROP TABLE IF EXISTS cbam_production_processes CASCADE"))
        db.flush()

    def upgrade() -> None:
        db.execute(
            text(
                """
CREATE TABLE IF NOT EXISTS cbam_production_processes (
    id UUID PRIMARY KEY,
    organization_id UUID NOT NULL REFERENCES organizations (id) ON DELETE CASCADE,
    reporting_period_binding_id UUID NOT NULL
        REFERENCES cbam_reporting_period_bindings (id) ON DELETE RESTRICT,
    installation_profile_id UUID NOT NULL
        REFERENCES cbam_installation_profiles (id) ON DELETE RESTRICT,
    product_profile_version_id UUID
        REFERENCES cbam_product_profile_versions (id) ON DELETE RESTRICT,
    name VARCHAR(255),
    identifier VARCHAR(128),
    calculation_method VARCHAR(64) NOT NULL DEFAULT 'CONVENTIONAL',
    status VARCHAR(32) NOT NULL DEFAULT 'draft',
    produced_quantity NUMERIC(24, 8),
    produced_quantity_unit VARCHAR(32),
    marketed_quantity NUMERIC(24, 8),
    marketed_quantity_unit VARCHAR(32),
    non_cbam_quantity NUMERIC(24, 8),
    non_cbam_quantity_unit VARCHAR(32),
    has_measurable_heat BOOLEAN,
    heat_imported_quantity NUMERIC(24, 8),
    heat_imported_unit VARCHAR(32),
    heat_exported_quantity NUMERIC(24, 8),
    heat_exported_unit VARCHAR(32),
    heat_imported_ef NUMERIC(24, 8),
    heat_exported_ef NUMERIC(24, 8),
    heat_ef_unit VARCHAR(32),
    heat_factor_source VARCHAR(255),
    heat_factor_document TEXT,
    has_waste_gas BOOLEAN,
    waste_gas_imported_quantity NUMERIC(24, 8),
    waste_gas_imported_unit VARCHAR(32),
    waste_gas_exported_quantity NUMERIC(24, 8),
    waste_gas_exported_unit VARCHAR(32),
    waste_gas_provenance TEXT,
    data_quality_code VARCHAR(128),
    data_verification_code VARCHAR(128),
    data_quality_justification_code VARCHAR(128),
    notes TEXT,
    row_version INTEGER NOT NULL DEFAULT 1,
    created_by_user_id UUID REFERENCES users (id) ON DELETE SET NULL,
    updated_by_user_id UUID REFERENCES users (id) ON DELETE SET NULL,
    created_at TIMESTAMP WITH TIME ZONE DEFAULT now() NOT NULL,
    updated_at TIMESTAMP WITH TIME ZONE DEFAULT now() NOT NULL
)
"""
            )
        )
        db.execute(
            text(
                """
CREATE TABLE IF NOT EXISTS cbam_production_process_product_uses (
    id UUID PRIMARY KEY,
    organization_id UUID NOT NULL REFERENCES organizations (id) ON DELETE CASCADE,
    reporting_period_binding_id UUID NOT NULL
        REFERENCES cbam_reporting_period_bindings (id) ON DELETE RESTRICT,
    process_id UUID NOT NULL REFERENCES cbam_production_processes (id) ON DELETE CASCADE,
    target_product_profile_version_id UUID NOT NULL
        REFERENCES cbam_product_profile_versions (id) ON DELETE RESTRICT,
    quantity NUMERIC(24, 8) NOT NULL,
    unit VARCHAR(32) NOT NULL,
    notes TEXT,
    row_version INTEGER NOT NULL DEFAULT 1,
    created_by_user_id UUID REFERENCES users (id) ON DELETE SET NULL,
    updated_by_user_id UUID REFERENCES users (id) ON DELETE SET NULL,
    created_at TIMESTAMP WITH TIME ZONE DEFAULT now() NOT NULL,
    updated_at TIMESTAMP WITH TIME ZONE DEFAULT now() NOT NULL
)
"""
            )
        )
        db.flush()

    def table_names() -> set[str]:
        return set(
            db.execute(
                text(
                    """
SELECT tablename FROM pg_tables
WHERE schemaname = 'public'
  AND tablename IN (
    'cbam_production_processes',
    'cbam_production_process_product_uses'
  )
"""
                )
            ).scalars()
        )

    downgrade()
    assert table_names() == set()
    upgrade()
    assert "cbam_production_processes" in table_names()
    downgrade()
    assert table_names() == set()
    upgrade()
    assert "cbam_production_process_product_uses" in table_names()
