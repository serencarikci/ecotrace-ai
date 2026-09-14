"""Phase 8A purchased-electricity math + workbook formula verification."""

from __future__ import annotations

from decimal import Decimal

import openpyxl
import pytest

from ecotrace.modules.cbam.application.purchased_electricity_constants import (
    WORKBOOK_ALLOCATION_FILENAME,
    WORKBOOK_ALLOCATION_SHA256,
    WORKBOOK_SEE_FILENAME,
    WORKBOOK_SEE_SHA256,
    local_reference_path,
)
from ecotrace.modules.cbam.application.purchased_electricity_math import (
    calculate_indirect_electricity_emissions,
    to_mwh,
)

GOLDEN_MWH = Decimal("100")
GOLDEN_FACTOR = Decimal("0.439")
GOLDEN_RESULT = Decimal("43.90000000")

# Facility-level SKDM sum(kWh)/1000 × H4 (no E/D attribution)
SKDM_TOTAL_KWH = Decimal("176034.53") + Decimal("180962.85") + Decimal("144768.9")
SKDM_TOTAL_MWH = SKDM_TOTAL_KWH / Decimal("1000")
SKDM_FACILITY_INDIRECT = SKDM_TOTAL_MWH * Decimal("0.439")

_ALLOC = local_reference_path(WORKBOOK_ALLOCATION_FILENAME)
_SEE = local_reference_path(WORKBOOK_SEE_FILENAME)
_HAS_WORKBOOKS = _ALLOC.is_file() and _SEE.is_file()


@pytest.mark.skipif(not _HAS_WORKBOOKS, reason="CBAM reference workbooks not in local-reference")
def test_workbook_files_present_and_hash() -> None:
    import hashlib

    alloc = _ALLOC
    see = _SEE
    assert alloc.is_file(), alloc
    assert see.is_file(), see
    assert hashlib.sha256(alloc.read_bytes()).hexdigest() == WORKBOOK_ALLOCATION_SHA256
    assert hashlib.sha256(see.read_bytes()).hexdigest() == WORKBOOK_SEE_SHA256


@pytest.mark.skipif(not _ALLOC.is_file(), reason="SKDM allocation workbook not in local-reference")
def test_workbook_skdm_electricity_cells_and_formula() -> None:
    path = local_reference_path(WORKBOOK_ALLOCATION_FILENAME)
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    assert ws["B2"].value == "Elektrik (KWh)"
    assert ws["G4"].value == "Elektrik EF (tCO2/MWh)"
    assert ws["H4"].value == pytest.approx(0.439)
    assert ws["C8"].value == "ELEKTRİK (MWh)"
    assert ws["D8"].value == "EF ELEKTRİK"
    assert ws["E8"].value == "ELEKTRİK TCO2E"
    assert ws["C9"].value == "=IF(D3=0,0,(E3/D3)*(B3/1000))"
    assert ws["D9"].value == "=$H$4"
    assert ws["E9"].value == "=D9*C9"
    assert ws["E12"].value == "=C12*D12"
    assert ws["B16"].value == "=$C$12"
    # Product electricity allocation cells exist but are out of Phase 8A
    assert str(ws["D19"].value).startswith("=IF(B19")


@pytest.mark.skipif(not _SEE.is_file(), reason="Official SEE workbook not in local-reference")
def test_workbook_see_electricity_formula_and_export_separate() -> None:
    path = local_reference_path(WORKBOOK_SEE_FILENAME)
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["D_Processes"]
    assert ws["L65"].value == pytest.approx(15.6817)
    assert ws["L66"].value == pytest.approx(0.439)
    assert ws["T66"].value == '=IF(G11="","",SUM(L65)*SUM(L66))'
    assert ws["L71"].value == 0
    assert ws["T72"].value == '=IF(G11="","",-SUM(L71)*SUM(L72))'
    # Export is a separate path; not subtracted from L65 in T66
    assert "L71" not in str(ws["T66"].value)


def test_mwh_times_compatible_factor_golden() -> None:
    out = calculate_indirect_electricity_emissions(
        electricity_quantity=GOLDEN_MWH,
        electricity_unit="MWh",
        factor_value=GOLDEN_FACTOR,
        factor_unit="tCO2e/MWh",
    )
    assert out.ok
    assert out.indirect_emissions_tco2e == GOLDEN_RESULT
    assert out.result_unit == "tCO2e"


def test_kwh_input_converts_to_mwh() -> None:
    out = calculate_indirect_electricity_emissions(
        electricity_quantity=Decimal("1000"),
        electricity_unit="kWh",
        factor_value=GOLDEN_FACTOR,
        factor_unit="tCO2e/MWh",
    )
    assert out.ok
    assert out.electricity_mwh == Decimal("1")
    assert out.indirect_emissions_tco2e == Decimal("0.43900000")


def test_skdm_facility_level_golden_without_ed_attribution() -> None:
    out = calculate_indirect_electricity_emissions(
        electricity_quantity=SKDM_TOTAL_KWH,
        electricity_unit="kWh",
        factor_value=Decimal("0.439"),
        factor_unit="tCO2/MWh",
    )
    assert out.ok
    assert out.electricity_mwh == SKDM_TOTAL_MWH
    # Quantized to 8 dp
    assert out.indirect_emissions_tco2e == Decimal("220.27539692")


def test_decimal_precision_no_float() -> None:
    out = calculate_indirect_electricity_emissions(
        electricity_quantity=Decimal("0.1"),
        electricity_unit="MWh",
        factor_value=Decimal("0.1"),
        factor_unit="tCO2e/MWh",
    )
    assert out.ok
    assert out.indirect_emissions_tco2e == Decimal("0.01000000")
    assert isinstance(out.indirect_emissions_tco2e, Decimal)


def test_zero_electricity_valid() -> None:
    out = calculate_indirect_electricity_emissions(
        electricity_quantity=Decimal("0"),
        electricity_unit="MWh",
        factor_value=GOLDEN_FACTOR,
        factor_unit="tCO2e/MWh",
    )
    assert out.ok
    assert out.indirect_emissions_tco2e == Decimal("0.00000000")


def test_zero_factor_valid_when_explicit() -> None:
    out = calculate_indirect_electricity_emissions(
        electricity_quantity=GOLDEN_MWH,
        electricity_unit="MWh",
        factor_value=Decimal("0"),
        factor_unit="tCO2e/MWh",
    )
    assert out.ok
    assert out.indirect_emissions_tco2e == Decimal("0.00000000")


def test_negative_electricity_rejected() -> None:
    out = calculate_indirect_electricity_emissions(
        electricity_quantity=Decimal("-1"),
        electricity_unit="MWh",
        factor_value=GOLDEN_FACTOR,
        factor_unit="tCO2e/MWh",
    )
    assert not out.ok
    assert out.error_code == "NEGATIVE_ELECTRICITY_QUANTITY"


def test_incompatible_units_rejected() -> None:
    out = calculate_indirect_electricity_emissions(
        electricity_quantity=Decimal("10"),
        electricity_unit="t",
        factor_value=GOLDEN_FACTOR,
        factor_unit="tCO2e/MWh",
    )
    assert not out.ok
    assert out.error_code == "INCOMPATIBLE_ELECTRICITY_UNIT"

    out2 = calculate_indirect_electricity_emissions(
        electricity_quantity=Decimal("10"),
        electricity_unit="MWh",
        factor_value=GOLDEN_FACTOR,
        factor_unit="tCO2e/t",
    )
    assert not out2.ok
    assert out2.error_code == "INCOMPATIBLE_FACTOR_UNIT"

    with pytest.raises(ValueError, match="INCOMPATIBLE_ELECTRICITY_UNIT"):
        to_mwh(Decimal("1"), "GJ")


def test_exported_not_part_of_math_helper() -> None:
    """Exported electricity is validated separately; calc ignores it."""
    out = calculate_indirect_electricity_emissions(
        electricity_quantity=Decimal("10"),
        electricity_unit="MWh",
        factor_value=Decimal("0.4"),
        factor_unit="tCO2e/MWh",
    )
    assert out.ok
    assert out.indirect_emissions_tco2e == Decimal("4.00000000")
