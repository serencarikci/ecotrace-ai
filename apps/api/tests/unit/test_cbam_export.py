from __future__ import annotations

import json
import uuid
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from tests.cbam_profile_helpers import create_active_ready_profile

from ecotrace.db.seed import DEMO_ORG_SLUG
from ecotrace.modules.cbam.application import (
    activity_record_service,
    allocation_rule_service,
    allocation_service,
    calculation_service,
    export_readiness_service,
    export_template_service,
    factor_catalog_service,
    factor_resolution_service,
    installation_service,
    period_binding_service,
    period_summary_service,
    production_record_service,
    workbook_export_service,
)
from ecotrace.modules.cbam.application.activity_record_service import ActivityRecordCreate
from ecotrace.modules.cbam.application.allocation_rule_service import (
    AllocationRuleCreate,
    AllocationRuleVersionRequest,
)
from ecotrace.modules.cbam.application.calculation_service import CalculationRunExecuteRequest
from ecotrace.modules.cbam.application.export_storage import sha256_file
from ecotrace.modules.cbam.application.factor_catalog_service import (
    FactorValueCreate,
    FactorValueVersionRequest,
)
from ecotrace.modules.cbam.application.installation_service import InstallationCreate
from ecotrace.modules.cbam.application.period_binding_service import (
    PeriodBindingCreate,
    PeriodBindingVersionRequest,
)
from ecotrace.modules.cbam.application.production_record_service import ProductionRecordCreate
from ecotrace.modules.cbam.architecture_boundary import find_forbidden_imports_in_tree
from ecotrace.modules.cbam.infrastructure.models import (
    CbamCalculationDefinition,
    CbamFactorDefinition,
    CbamReferenceSource,
)
from ecotrace.modules.facilities.infrastructure.models import Facility
from ecotrace.modules.identity.infrastructure.models import User
from ecotrace.modules.organizations.infrastructure.models import Organization
from ecotrace.modules.reporting_periods.infrastructure.models import ReportingPeriod


def _org(db):
    return db.execute(select(Organization).where(Organization.slug == DEMO_ORG_SLUG)).scalar_one()


def _admin(db):
    return db.execute(
        select(User).where(User.normalized_email == 'orgadmin@ecotrace.dev')
    ).scalar_one()


def _facility(db, org_id):
    return db.execute(select(Facility).where(Facility.organization_id == org_id).limit(1)).scalar_one()


def _setup(db, admin, org):
    facility = _facility(db, org.id)
    installation_service.create_installation(
        db,
        admin,
        org.id,
        InstallationCreate(
            facility_id=facility.id,
            code=f'P6-{uuid.uuid4().hex[:8]}',
            name='P6 Installation',
        ),
    )
    period = ReportingPeriod(
        organization_id=org.id,
        code=f'P6-{uuid.uuid4().hex[:6]}',
        name='P6 Period',
        period_type='custom',
        start_date=date(2034, 1, 1),
        end_date=date(2034, 3, 31),
        status='open',
    )
    db.add(period)
    db.commit()
    binding = period_binding_service.create_period_binding(
        db, admin, org.id, PeriodBindingCreate(reporting_period_id=period.id)
    )
    opened = period_binding_service.open_data_collection(
        db,
        admin,
        org.id,
        binding.id,
        PeriodBindingVersionRequest(row_version=binding.row_version),
    )
    installation = installation_service.list_installations(
        db, admin, org.id, page=1, page_size=50
    ).items[-1]
    return opened, installation


def _definition(db, code: str) -> CbamFactorDefinition:
    return db.execute(
        select(CbamFactorDefinition).where(CbamFactorDefinition.code == code)
    ).scalar_one()


def _manual_source(db) -> CbamReferenceSource:
    return db.execute(
        select(CbamReferenceSource).where(
            CbamReferenceSource.organization_id.is_(None),
            CbamReferenceSource.code == 'MANUAL_APPROVED_REFERENCE',
        )
    ).scalar_one()


def _activate_value(db, admin, org_id, definition_id, **kwargs):
    created = factor_catalog_service.create_factor_value(
        db,
        admin,
        org_id,
        definition_id,
        FactorValueCreate(
            reference_source_id=_manual_source(db).id,
            **kwargs,
        ),
    )
    return factor_catalog_service.activate_factor_value(
        db,
        admin,
        org_id,
        created.id,
        FactorValueVersionRequest(row_version=created.row_version),
    )


def _prepare_exportable_scenario(db, admin, org):
    binding, installation = _setup(db, admin, org)
    profile = create_active_ready_profile(db, admin, org.id)
    base = production_record_service.create_production_record(
        db,
        admin,
        org.id,
        binding.id,
        ProductionRecordCreate(
            product_profile_version_id=profile.id,
            installation_profile_id=installation.id,
            quantity=Decimal('500'),
            unit='t',
        ),
    )
    target = production_record_service.create_production_record(
        db,
        admin,
        org.id,
        binding.id,
        ProductionRecordCreate(
            product_profile_version_id=profile.id,
            installation_profile_id=installation.id,
            quantity=Decimal('100'),
            unit='t',
        ),
    )
    activity = activity_record_service.create_activity_record(
        db,
        admin,
        org.id,
        binding.id,
        ActivityRecordCreate(
            installation_profile_id=installation.id,
            activity_type='ELECTRICITY',
            quantity=Decimal('100'),
            unit='MWh',
            data_source_type='PRIMARY',
            activity_date=date(2034, 2, 1),
        ),
    )
    rule = allocation_rule_service.create_allocation_rule(
        db,
        admin,
        org.id,
        binding.id,
        AllocationRuleCreate(
            name='P6 ratio',
            installation_profile_id=installation.id,
            allocation_method='PRODUCTION_QUANTITY_RATIO',
            numerator_production_record_id=target.id,
            denominator_production_record_id=base.id,
        ),
    )
    active = allocation_rule_service.activate_allocation_rule(
        db,
        admin,
        org.id,
        rule.id,
        AllocationRuleVersionRequest(row_version=rule.row_version),
    )
    alloc = allocation_service.allocate_activity_record(
        db, admin, org.id, active.id, activity.id
    )
    assert alloc.allocated_quantity == Decimal('20.00000000')

    ef = _definition(db, 'GENERIC_EMISSION_FACTOR')
    _activate_value(
        db,
        admin,
        org.id,
        ef.id,
        activity_type='ELECTRICITY',
        numeric_value=Decimal('0.4'),
        unit='tCO2e/MWh',
        data_source_type='DEFAULT_REFERENCE',
        valid_from=date(2030, 1, 1),
        valid_until=date(2040, 12, 31),
    )
    factor_resolution_service.resolve_for_allocation_result(
        db, admin, org.id, alloc.id, 'GENERIC_EMISSION_FACTOR'
    )
    calc_def = db.execute(
        select(CbamCalculationDefinition).where(
            CbamCalculationDefinition.code == 'MULTIPLY_ALLOCATION_BY_GENERIC_EF'
        )
    ).scalar_one()
    run = calculation_service.create_calculation_run(db, admin, org.id, binding.id)
    executed = calculation_service.execute_calculation_run(
        db,
        admin,
        org.id,
        run.id,
        CalculationRunExecuteRequest(calculation_definition_ids=[calc_def.id]),
    )
    results = calculation_service.list_calculation_results(
        db, admin, org.id, executed.id, page=1, page_size=50
    ).items
    row = next(r for r in results if r.source_type == 'ALLOCATION_RESULT')
    assert row.status == 'CALCULATED'
    assert row.result_value == Decimal('8.00000000')
    assert row.source_quantity == Decimal('20.00000000')
    assert row.factor_value == Decimal('0.40000000')
    return binding, installation, executed, alloc


def test_internal_template_loads_and_maps(seeded_db) -> None:
    org = _org(seeded_db)
    admin = _admin(seeded_db)
    template = export_template_service.ensure_internal_export_template(seeded_db)
    path = export_template_service.verify_template_file(template)
    assert path.is_file()
    original_checksum = sha256_file(path)
    wb = load_workbook(path)
    assert set(wb.sheetnames) >= {
        'Organization',
        'Installation',
        'Production',
        'Activities',
        'Purchased Inputs',
        'Allocation',
        'Factors',
        'Calculations',
        'Summary',
    }
    assert wb['Summary']['E3'].value == '=COUNTA(A4:A18)'
    assert 'INTERNAL DEVELOPMENT TEMPLATE' in str(wb['Organization']['A1'].value)

    binding, _installation, executed, _alloc = _prepare_exportable_scenario(
        seeded_db, admin, org
    )
    readiness = export_readiness_service.assess_export_readiness(
        seeded_db, admin, org.id, binding.id, calculation_run_id=executed.id
    )
    assert readiness.status in {'READY', 'READY_WITH_WARNINGS'}
    assert readiness.official_mapping_blocked is True

    export_run = workbook_export_service.create_export_run(
        seeded_db,
        admin,
        org.id,
        binding.id,
        workbook_export_service.ExportCreateRequest(calculation_run_id=executed.id),
    )
    assert export_run.status in {'COMPLETED', 'COMPLETED_WITH_WARNINGS'}
    assert sha256_file(path) == original_checksum

    artifacts = workbook_export_service.list_export_artifacts(
        seeded_db, admin, org.id, export_run.id
    )
    xlsx = next(a for a in artifacts if a.artifact_type == 'XLSX')
    manifest = next(a for a in artifacts if a.file_name == 'export-manifest.json')
    out_path = workbook_export_service.download_export_artifact(
        seeded_db, admin, org.id, xlsx.id
    )[1]
    assert out_path.is_file()
    reopened = load_workbook(out_path)
    assert reopened['Summary']['E3'].value == '=COUNTA(A4:A18)'
    assert reopened['Organization']['B3'].value == org.name
    calc_sheet = reopened['Calculations']
    assert calc_sheet['C3'].value == 20.0
    assert calc_sheet['E3'].value == 0.4
    assert calc_sheet['G3'].value == 8.0
    alloc_sheet = reopened['Allocation']
    assert alloc_sheet['E3'].value == 20.0

    manifest_path = workbook_export_service.download_export_artifact(
        seeded_db, admin, org.id, manifest.id
    )[1]
    payload = json.loads(manifest_path.read_text(encoding='utf-8'))
    assert payload['reportingPeriodBindingId'] == str(binding.id)
    assert payload['calculationRunId'] == str(executed.id)
    assert payload['templateChecksum'] == template.checksum
    assert payload['mappingVersion'] == template.mapping_version
    assert payload['officialMappingBlocked'] is True
    assert payload['traceability']['calculationResultIds']
    assert payload['traceability']['allocationResultIds']
    assert payload['traceability']['factorResolutionIds']
    assert payload['artifactChecksum'] == xlsx.sha256


def test_readiness_states(seeded_db) -> None:
    org = _org(seeded_db)
    admin = _admin(seeded_db)
    binding, installation = _setup(seeded_db, admin, org)

    not_ready = export_readiness_service.assess_export_readiness(
        seeded_db, admin, org.id, binding.id
    )
    assert not_ready.status == 'NOT_READY'
    assert any(c.code == 'calculation' and c.status == 'MISSING' for c in not_ready.checks)

    activity_record_service.create_activity_record(
        seeded_db,
        admin,
        org.id,
        binding.id,
        ActivityRecordCreate(
            installation_profile_id=installation.id,
            activity_type='ELECTRICITY',
            quantity=Decimal('10'),
            unit='kWh',
            data_source_type='PRIMARY',
            activity_date=date(2034, 2, 1),
        ),
    )
    still_not_ready = export_readiness_service.assess_export_readiness(
        seeded_db, admin, org.id, binding.id
    )
    assert still_not_ready.status == 'NOT_READY'

    binding2, _inst2, executed, _alloc = _prepare_exportable_scenario(seeded_db, admin, org)
    ready = export_readiness_service.assess_export_readiness(
        seeded_db, admin, org.id, binding2.id, calculation_run_id=executed.id
    )
    assert ready.status in {'READY', 'READY_WITH_WARNINGS'}
    assert ready.official_mapping_blocked is True
    assert ready.status == 'READY_WITH_WARNINGS'


def test_blocked_calculation_not_exported_as_zero(seeded_db) -> None:
    org = _org(seeded_db)
    admin = _admin(seeded_db)
    binding, installation = _setup(seeded_db, admin, org)
    ef = _definition(seeded_db, 'GENERIC_EMISSION_FACTOR')
    activity = activity_record_service.create_activity_record(
        seeded_db,
        admin,
        org.id,
        binding.id,
        ActivityRecordCreate(
            installation_profile_id=installation.id,
            activity_type='ELECTRICITY',
            quantity=Decimal('10'),
            unit='kWh',
            data_source_type='PRIMARY',
            activity_date=date(2034, 2, 1),
        ),
    )
    _activate_value(
        seeded_db,
        admin,
        org.id,
        ef.id,
        activity_type='ELECTRICITY',
        numeric_value=Decimal('0.4'),
        unit='kgCO2e/kWh',
        data_source_type='DEFAULT_REFERENCE',
        valid_from=date(2030, 1, 1),
        valid_until=date(2040, 12, 31),
    )
    _activate_value(
        seeded_db,
        admin,
        org.id,
        ef.id,
        activity_type='ELECTRICITY',
        numeric_value=Decimal('0.5'),
        unit='kgCO2e/kWh',
        data_source_type='DEFAULT_REFERENCE',
        valid_from=date(2030, 1, 1),
        valid_until=date(2040, 12, 31),
    )
    factor_resolution_service.resolve_for_activity_record(
        seeded_db, admin, org.id, activity.id, 'GENERIC_EMISSION_FACTOR'
    )
    run = calculation_service.create_calculation_run(seeded_db, admin, org.id, binding.id)
    executed = calculation_service.execute_calculation_run(
        seeded_db,
        admin,
        org.id,
        run.id,
        CalculationRunExecuteRequest(allow_unallocated_activity=True),
    )
    results = calculation_service.list_calculation_results(
        seeded_db, admin, org.id, executed.id, page=1, page_size=50
    ).items
    blocked = [r for r in results if r.source_id == activity.id]
    assert blocked
    assert all(r.result_value is None for r in blocked)

    readiness = export_readiness_service.assess_export_readiness(
        seeded_db, admin, org.id, binding.id, calculation_run_id=executed.id
    )
    assert readiness.status == 'NOT_READY'
    calc_check = next(c for c in readiness.checks if c.code == 'calculation')
    assert calc_check.status == 'MISSING'
    results = calculation_service.list_calculation_results(
        seeded_db, admin, org.id, executed.id, page=1, page_size=50
    ).items
    activity_results = [r for r in results if r.source_id == activity.id]
    assert activity_results
    assert all(r.result_value is None for r in activity_results)
    assert all(r.result_value != Decimal('0') for r in activity_results)


def test_period_summary_title(seeded_db) -> None:
    org = _org(seeded_db)
    admin = _admin(seeded_db)
    binding, _installation = _setup(seeded_db, admin, org)
    summary = period_summary_service.get_period_summary(seeded_db, admin, org.id, binding.id)
    assert summary.title == 'SKDM Period Summary'
    assert summary.official_mapping_blocked is True
    assert 'not an official' in summary.disclaimer.lower()


def test_excel_formula_injection_is_neutralized() -> None:
    from ecotrace.modules.cbam.application.workbook_export_service import (
        _excel_safe_cell_value,
    )

    assert _excel_safe_cell_value('=1+1') == "'=1+1"
    assert _excel_safe_cell_value('+cmd') == "'+cmd"
    assert _excel_safe_cell_value('-2+3') == "'-2+3"
    assert _excel_safe_cell_value('@SUM(A1)') == "'@SUM(A1)"
    assert _excel_safe_cell_value('normal text') == 'normal text'
    assert _excel_safe_cell_value(8.0) == 8.0
    assert _excel_safe_cell_value(None) is None


def test_phase6_does_not_become_second_calculation_engine() -> None:
    assert find_forbidden_imports_in_tree() == []
    api_root = Path(__file__).resolve().parents[2]
    export_paths = [
        api_root / 'src/ecotrace/modules/cbam/application/workbook_export_service.py',
        api_root / 'src/ecotrace/modules/cbam/application/export_context.py',
        api_root / 'src/ecotrace/modules/cbam/application/export_readiness_service.py',
        api_root / 'src/ecotrace/modules/cbam/application/period_summary_service.py',
        api_root / 'src/ecotrace/modules/cbam/application/internal_template_builder.py',
    ]
    banned = (
        'multiply_activity_by_factor',
        'source_quantity * allocation_ratio',
        'fetch_ipcc',
        'fetch_defra',
        'gwp_table',
        'certificate_quantity',
        'financial_obligation',
        'eval(',
    )
    for path in export_paths:
        text = path.read_text(encoding='utf-8').lower()
        for token in banned:
            assert token not in text, f'{path} contains {token}'
        assert 'activity_quantity *' not in text
        assert 'allocated_quantity *' not in text
