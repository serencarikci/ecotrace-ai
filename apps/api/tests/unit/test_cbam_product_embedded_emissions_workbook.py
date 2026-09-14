"""Phase 10C workbook provenance: D_Processes T54/T58/T62/T66/T72 and Summary_Products."""

from __future__ import annotations

import hashlib
import os
from pathlib import Path

import pytest

from ecotrace.modules.cbam.application.product_embedded_emissions_constants import (
    WORKBOOK_FILENAME,
    WORKBOOK_PRIMARY_SHEET,
    WORKBOOK_SHA256,
    WORKBOOK_SUMMARY_SHEET,
)

WORKBOOK = Path(
    os.environ.get(
        'CBAM_SEE_WORKBOOK_PATH',
        str(
            Path(__file__).resolve().parents[4] / 'local-reference' / WORKBOOK_FILENAME
        ),
    )
)

pytestmark = pytest.mark.skipif(
    not WORKBOOK.is_file(), reason='CBAM SEE workbook not available'
)


def test_workbook_checksum_matches_constants() -> None:
    assert hashlib.sha256(WORKBOOK.read_bytes()).hexdigest() == WORKBOOK_SHA256


def test_d_processes_attributed_emission_formulas_and_signs() -> None:
    from openpyxl import load_workbook

    wb = load_workbook(WORKBOOK, data_only=False)
    ws = wb[WORKBOOK_PRIMARY_SHEET]

    # Denominator: produced quantity of the process (TotProd_), not marketed.
    assert ws['L24'].value == '=IF(G11="","",SUM(L16:L23))'

    # T54 mirrors the allocated direct emissions of the process (DEA product row).
    assert ws['T54'].value == '=IF(G11="","",L54)'

    # T58 measurable heat: imported minus exported, both signed positive/negative.
    assert ws['T58'].value == '=IF(G11="","",SUM(L57*L58-M57*M58))'

    # T62 waste gas: import at EF, export credited at EF*0.667 (negative sign).
    assert ws['T62'].value == (
        '=IF(G11="","",SUM(L61*CONST_EFNatGas-M61*CONST_EFNatGas*0.667))'
    )

    # T66 indirect: electricity consumption × electricity emission factor (positive).
    assert ws['T66'].value == '=IF(G11="","",SUM(L65)*SUM(L66))'

    # T72 exported electricity: negative sign, and it feeds the DIRECT bucket only.
    assert ws['T72'].value == '=IF(G11="","",-SUM(L71)*SUM(L72))'
    assert ws['S72'].value == '=CONST_CNTR_EmbedEmDir&G11'


def test_summary_products_specific_columns_are_direct_indirect_and_their_sum() -> None:
    from openpyxl import load_workbook

    wb = load_workbook(WORKBOOK, data_only=False)
    ws = wb[WORKBOOK_SUMMARY_SHEET]

    assert ws['I10'].value == (
        '=IF(OR($D10="",$D10=CONST_NA),"",'
        'INDEX(InputOutput!$AK$71:$AK$80,MATCH($D10,InputOutput!$D$71:$D$80,0)))'
    )
    assert ws['J10'].value == (
        '=IF(OR($D10="",$D10=CONST_NA),"",'
        'INDEX(InputOutput!$AM$71:$AM$80,MATCH($D10,InputOutput!$D$71:$D$80,0)))'
    )
    # K = I + J is exactly the specific_total identity implemented in the math module.
    assert ws['K10'].value == '=IF(OR($D10="",$D10=CONST_NA),"",SUM(I10:J10))'


def test_workbook_product_1_reference_values_document_the_v1_leontief_gap() -> None:
    """Workbook Product 1 uses the full Leontief inverse over four precursors.

    EcoTrace V1 models purchased precursors only (no process-as-precursor), so these
    numbers are reference-only and are not reproducible by the roll-up engine.
    """
    from openpyxl import load_workbook

    wb = load_workbook(WORKBOOK, data_only=True)
    ws = wb[WORKBOOK_SUMMARY_SHEET]

    direct, indirect, total = ws['I10'].value, ws['J10'].value, ws['K10'].value
    assert direct == pytest.approx(0.868957674443175)
    assert indirect == pytest.approx(0.37707768433481187)
    assert total == pytest.approx(direct + indirect)
    assert ws['L10'].value == 'tCO2e/t'
