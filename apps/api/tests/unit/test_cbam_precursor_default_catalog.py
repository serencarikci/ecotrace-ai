"""Phase 10A EU precursor default-value catalog: seed, immutability, resolution.

The 12 532-row catalog is deliberately NOT seeded by the global conftest. Every test
here lazy-seeds it through ``ensure_platform_precursor_default_catalog``.
"""

from __future__ import annotations

import hashlib
import os
from decimal import Decimal
from pathlib import Path

import pytest
from sqlalchemy import func, select, text

from ecotrace.core.exceptions import ConflictError
from ecotrace.modules.cbam.application import precursor_default_catalog_seed
from ecotrace.modules.cbam.application.precursor_constants import (
    CODE_DEFAULT_VALUE_AMBIGUOUS,
    CODE_DEFAULT_VALUE_UNRESOLVED,
    CODE_PRECURSOR_CN_CODE_REQUIRED,
    CODE_PRECURSOR_COUNTRY_REQUIRED,
    DEFAULT_DATASET_CODE,
    DEFAULT_DATASET_STATUS_ACTIVE,
    DEFAULT_DATASET_VERSION,
    DV_STATUS_DASH,
    DV_STATUS_NUMERIC,
    OTHER_COUNTRIES_GROUP_NAME,
    RESOLUTION_AMBIGUOUS,
    RESOLUTION_RESOLVED,
    RESOLUTION_UNRESOLVED,
    SPECIFIC_DIRECT_UNIT,
    SPECIFIC_INDIRECT_UNIT,
    WORKBOOK_PRIMARY_SHEET,
    WORKBOOK_SHA256,
)
from ecotrace.modules.cbam.application.precursor_default_catalog_seed import (
    build_lookup_key,
    compute_seed_content_checksum,
    ensure_platform_precursor_default_catalog,
    load_precursor_default_seed_payload,
    normalize_lookup_segment,
)
from ecotrace.modules.cbam.application.precursor_default_catalog_service import (
    build_default_snapshot,
    get_active_dataset,
    normalize_country_name,
    normalize_precursor_cn_code,
    resolve_default_value,
    snapshot_specific_values,
)
from ecotrace.modules.cbam.infrastructure.models import (
    CbamPrecursorDefaultDataset,
    CbamPrecursorDefaultValue,
)

EXPECTED_VALUE_COUNT = 12532
EXPECTED_CONTENT_CHECKSUM = "7543752e2ba7ccb314e6037355dfa2e11c65a0f129f6d27a0c5ffdd167c645c4"
DV_WORKBOOK_SHA256 = "865372ed23649b7b02c9124f207fc0b0875fd244c45c19e9fb8cdb1e503a5003"

_REPO_ROOT = Path(__file__).resolve().parents[4]
REFERENCE_DIR = Path(os.environ.get("CBAM_REFERENCE_DIR", str(_REPO_ROOT / "local-reference")))
SEE_WORKBOOK = REFERENCE_DIR / (
    "CBAM SEE V2.1_Example Steel 3 Screws and nuts_final Dosyasının Kopyası- (1) (1).xlsx"
)
DV_WORKBOOK = REFERENCE_DIR / "DVs_as_adopted_v20260204.xlsx"

# Deterministic fixtures picked from the seed payload.
UNIQUE_COUNTRY = "Albania"
UNIQUE_CN_DISPLAY = "2523 29 00"
UNIQUE_CN_NORMALIZED = "25232900"
UNIQUE_DIRECT = Decimal("0.9")
UNIQUE_INDIRECT = Decimal("0.03")

# Albania 2523 90 00 exists only per production route, never with a NULL route.
ROUTED_CN = "25239000"
ROUTED_ROUTE = "(A)"
ROUTED_DIRECT = Decimal("0.86")

# Argentina 2523 90 00 has grey + white hydraulic cements on a NULL route.
AMBIGUOUS_COUNTRY = "Argentina"
AMBIGUOUS_CN = "25239000"
AMBIGUOUS_GREY = "Grey hydraulic cements"
AMBIGUOUS_WHITE = "White hydraulic cements"

# 2523 10 00 route (B) is published only for the Other Countries group.
GROUP_ONLY_CN = "25231000"
GROUP_ONLY_ROUTE = "(B)"


# --------------------------------------------------------------------------------------
# Workbook provenance (skipped when the reference workbooks are not checked out)
# --------------------------------------------------------------------------------------


@pytest.mark.skipif(not SEE_WORKBOOK.is_file(), reason="CBAM SEE workbook not available")
def test_see_e_purchprec_cells_and_formulas() -> None:
    """L25/L39/L52/T49/T52 are the formulas Phase 10A reimplements."""
    from openpyxl import load_workbook

    assert hashlib.sha256(SEE_WORKBOOK.read_bytes()).hexdigest() == WORKBOOK_SHA256
    wb = load_workbook(SEE_WORKBOOK, data_only=False)
    try:
        assert WORKBOOK_PRIMARY_SHEET in wb.sheetnames
        ws = wb[WORKBOOK_PRIMARY_SHEET]
        assert ws["L25"].value == '=IF(G14="","",SUM(L17:L24))'
        assert ws["L39"].value == '=IF(G14="","",SUM(L25)-SUM(L28:L38))'
        assert ws["L52"].value == '=IF(COUNT(L50:L51)=0,"",L50*L51)'
        assert "SUM(T25)*SUM(L49)" in str(ws["T49"].value)
        assert "SUM(T25)*SUM(L52)" in str(ws["T52"].value)
    finally:
        wb.close()


@pytest.mark.skipif(not DV_WORKBOOK.is_file(), reason="DV workbook not available")
def test_dv_workbook_checksum_and_extraction_readable() -> None:
    """The seed payload declares the exact DV workbook it was extracted from."""
    from openpyxl import load_workbook

    assert hashlib.sha256(DV_WORKBOOK.read_bytes()).hexdigest() == DV_WORKBOOK_SHA256
    payload = load_precursor_default_seed_payload()
    assert payload["dataset"]["sourceWorkbookSha256"] == DV_WORKBOOK_SHA256
    assert payload["dataset"]["sourceWorkbookName"] == DV_WORKBOOK.name

    wb = load_workbook(DV_WORKBOOK, read_only=True, data_only=True)
    try:
        sheets = set(wb.sheetnames)
        assert {"Overview", "Version History"} <= sheets
        seeded_sheets = {row["sourceSheet"] for row in payload["values"]}
        assert seeded_sheets <= sheets
    finally:
        wb.close()


# --------------------------------------------------------------------------------------
# Seed determinism and immutability
# --------------------------------------------------------------------------------------


def test_seed_payload_checksum_is_deterministic() -> None:
    payload = load_precursor_default_seed_payload()
    assert len(payload["values"]) == EXPECTED_VALUE_COUNT
    assert payload["dataset"]["valueCount"] == EXPECTED_VALUE_COUNT
    assert compute_seed_content_checksum(payload) == EXPECTED_CONTENT_CHECKSUM
    assert payload["dataset"]["contentChecksum"] == EXPECTED_CONTENT_CHECKSUM
    # The checksum must not depend on the stored value itself.
    without_checksum = load_precursor_default_seed_payload()
    without_checksum["dataset"].pop("contentChecksum")
    assert compute_seed_content_checksum(without_checksum) == EXPECTED_CONTENT_CHECKSUM


def test_seed_inserts_deterministic_dataset_and_values(seeded_db) -> None:
    dataset = ensure_platform_precursor_default_catalog(seeded_db)
    assert dataset.dataset_code == DEFAULT_DATASET_CODE
    assert dataset.dataset_version == DEFAULT_DATASET_VERSION
    assert dataset.status == DEFAULT_DATASET_STATUS_ACTIVE
    assert dataset.content_checksum == EXPECTED_CONTENT_CHECKSUM
    assert dataset.source_workbook_sha256 == DV_WORKBOOK_SHA256
    assert dataset.value_count == EXPECTED_VALUE_COUNT

    count = seeded_db.execute(
        select(func.count())
        .select_from(CbamPrecursorDefaultValue)
        .where(CbamPrecursorDefaultValue.dataset_id == dataset.id)
    ).scalar_one()
    assert count == EXPECTED_VALUE_COUNT

    row = seeded_db.execute(
        select(CbamPrecursorDefaultValue).where(
            CbamPrecursorDefaultValue.dataset_id == dataset.id,
            CbamPrecursorDefaultValue.country_name == UNIQUE_COUNTRY,
            CbamPrecursorDefaultValue.cn_normalized_code == UNIQUE_CN_NORMALIZED,
        )
    ).scalar_one()
    assert row.direct_value == UNIQUE_DIRECT
    assert row.indirect_value == UNIQUE_INDIRECT
    assert row.direct_value_status == DV_STATUS_NUMERIC
    assert row.lookup_key == build_lookup_key(
        country_name=UNIQUE_COUNTRY,
        cn_normalized_code=UNIQUE_CN_NORMALIZED,
        production_route=None,
        goods_description=row.goods_description,
    )


def test_seed_is_idempotent(seeded_db) -> None:
    first = ensure_platform_precursor_default_catalog(seeded_db)
    seeded_db.flush()
    second = ensure_platform_precursor_default_catalog(seeded_db)
    assert first.id == second.id

    datasets = seeded_db.execute(
        select(func.count()).select_from(CbamPrecursorDefaultDataset)
    ).scalar_one()
    values = seeded_db.execute(
        select(func.count()).select_from(CbamPrecursorDefaultValue)
    ).scalar_one()
    assert datasets == 1
    assert values == EXPECTED_VALUE_COUNT


def test_published_dataset_checksum_mutation_conflicts(seeded_db) -> None:
    dataset = ensure_platform_precursor_default_catalog(seeded_db)
    seeded_db.execute(
        text("UPDATE cbam_precursor_default_datasets SET content_checksum = :c WHERE id = :i"),
        {"c": "0" * 64, "i": dataset.id},
    )
    seeded_db.expire_all()

    with pytest.raises(ConflictError) as exc:
        ensure_platform_precursor_default_catalog(seeded_db)
    details = str(exc.value.details)
    assert "IMMUTABLE_PRECURSOR_DV_DATASET_CONFLICT" in details
    assert "content_checksum" in details


def test_published_dataset_row_count_drift_conflicts(seeded_db) -> None:
    dataset = ensure_platform_precursor_default_catalog(seeded_db)
    victim = seeded_db.execute(
        select(CbamPrecursorDefaultValue)
        .where(CbamPrecursorDefaultValue.dataset_id == dataset.id)
        .limit(1)
    ).scalar_one()
    seeded_db.delete(victim)
    seeded_db.flush()

    with pytest.raises(ConflictError) as exc:
        ensure_platform_precursor_default_catalog(seeded_db)
    assert "default_value_count" in str(exc.value.details)


def test_seed_payload_checksum_mismatch_is_rejected(seeded_db, monkeypatch) -> None:
    payload = load_precursor_default_seed_payload()
    payload["dataset"]["contentChecksum"] = "f" * 64
    monkeypatch.setattr(
        precursor_default_catalog_seed,
        "load_precursor_default_seed_payload",
        lambda: payload,
    )
    with pytest.raises(ConflictError) as exc:
        ensure_platform_precursor_default_catalog(seeded_db)
    assert "SEED_CHECKSUM_MISMATCH" in str(exc.value.details)


def test_get_active_dataset_lazy_seeds(seeded_db) -> None:
    assert (
        seeded_db.execute(
            select(func.count()).select_from(CbamPrecursorDefaultDataset)
        ).scalar_one()
        == 0
    )
    dataset = get_active_dataset(seeded_db)
    assert dataset.content_checksum == EXPECTED_CONTENT_CHECKSUM
    assert dataset.value_count == EXPECTED_VALUE_COUNT


# --------------------------------------------------------------------------------------
# Key normalization
# --------------------------------------------------------------------------------------


def test_normalizers_and_lookup_key() -> None:
    assert normalize_precursor_cn_code("2523 29 00") == "25232900"
    assert normalize_precursor_cn_code("ex 7318.15.95") == "73181595"
    assert normalize_country_name("  Albania\xa0 ") == "Albania"
    assert normalize_lookup_segment(None) == ""
    assert normalize_lookup_segment(" Grey  Portland\xa0Cement ") == "grey portland cement"
    assert (
        build_lookup_key(
            country_name="Albania",
            cn_normalized_code="25232900",
            production_route=None,
            goods_description=None,
        )
        == "albania|25232900||"
    )


# --------------------------------------------------------------------------------------
# Resolution
# --------------------------------------------------------------------------------------


def test_resolve_exact_default_is_resolved(seeded_db) -> None:
    ensure_platform_precursor_default_catalog(seeded_db)
    resolution = resolve_default_value(
        seeded_db,
        country_of_origin=UNIQUE_COUNTRY,
        cn_code=UNIQUE_CN_DISPLAY,
        production_route=None,
    )
    assert resolution.status == RESOLUTION_RESOLVED
    assert resolution.issue_code is None
    assert resolution.candidate_count == 1
    assert resolution.value is not None
    assert resolution.value.country_name == UNIQUE_COUNTRY
    assert resolution.value.cn_normalized_code == UNIQUE_CN_NORMALIZED
    assert resolution.value.direct_value == UNIQUE_DIRECT
    assert resolution.value.indirect_value == UNIQUE_INDIRECT
    assert resolution.value.direct_unit == SPECIFIC_DIRECT_UNIT
    assert resolution.value.indirect_unit == SPECIFIC_INDIRECT_UNIT
    assert resolution.dataset is not None
    assert resolution.dataset.content_checksum == EXPECTED_CONTENT_CHECKSUM


def test_resolve_unknown_cn_is_unresolved(seeded_db) -> None:
    ensure_platform_precursor_default_catalog(seeded_db)
    resolution = resolve_default_value(
        seeded_db, country_of_origin=UNIQUE_COUNTRY, cn_code="9999 99 99"
    )
    assert resolution.status == RESOLUTION_UNRESOLVED
    assert resolution.issue_code == CODE_DEFAULT_VALUE_UNRESOLVED
    assert resolution.value is None
    assert resolution.candidate_count == 0


def test_resolve_requires_country_and_cn(seeded_db) -> None:
    ensure_platform_precursor_default_catalog(seeded_db)
    no_country = resolve_default_value(seeded_db, country_of_origin=None, cn_code=UNIQUE_CN_DISPLAY)
    assert no_country.status == RESOLUTION_UNRESOLVED
    assert no_country.issue_code == CODE_PRECURSOR_COUNTRY_REQUIRED

    no_cn = resolve_default_value(seeded_db, country_of_origin=UNIQUE_COUNTRY, cn_code=None)
    assert no_cn.status == RESOLUTION_UNRESOLVED
    assert no_cn.issue_code == CODE_PRECURSOR_CN_CODE_REQUIRED
    assert CODE_PRECURSOR_CN_CODE_REQUIRED == "PRECURSOR_CN_CODE_REQUIRED"


def test_resolve_ambiguous_without_description(seeded_db) -> None:
    """Grey and white cement share CN + NULL route; only a description separates them."""
    ensure_platform_precursor_default_catalog(seeded_db)
    ambiguous = resolve_default_value(
        seeded_db, country_of_origin=AMBIGUOUS_COUNTRY, cn_code=AMBIGUOUS_CN
    )
    assert ambiguous.status == RESOLUTION_AMBIGUOUS
    assert ambiguous.issue_code == CODE_DEFAULT_VALUE_AMBIGUOUS
    assert ambiguous.value is None
    assert ambiguous.candidate_count == 2
    assert {c.goods_description for c in ambiguous.candidates} == {
        AMBIGUOUS_GREY,
        AMBIGUOUS_WHITE,
    }


def test_description_narrows_ambiguous_match(seeded_db) -> None:
    ensure_platform_precursor_default_catalog(seeded_db)
    narrowed = resolve_default_value(
        seeded_db,
        country_of_origin=AMBIGUOUS_COUNTRY,
        cn_code=AMBIGUOUS_CN,
        goods_description="  grey   hydraulic cements ",
    )
    assert narrowed.status == RESOLUTION_RESOLVED
    assert narrowed.value is not None
    assert narrowed.value.goods_description == AMBIGUOUS_GREY
    assert narrowed.value.direct_value_status == DV_STATUS_DASH


def test_null_route_matches_null_route_only(seeded_db) -> None:
    ensure_platform_precursor_default_catalog(seeded_db)
    without_route = resolve_default_value(
        seeded_db, country_of_origin=UNIQUE_COUNTRY, cn_code=ROUTED_CN
    )
    assert without_route.status == RESOLUTION_UNRESOLVED
    assert without_route.issue_code == CODE_DEFAULT_VALUE_UNRESOLVED

    with_route = resolve_default_value(
        seeded_db,
        country_of_origin=UNIQUE_COUNTRY,
        cn_code=ROUTED_CN,
        production_route=ROUTED_ROUTE,
    )
    assert with_route.status == RESOLUTION_RESOLVED
    assert with_route.value is not None
    assert with_route.value.production_route == ROUTED_ROUTE
    assert with_route.value.direct_value == ROUTED_DIRECT

    # Route matching is case/whitespace insensitive but still exact on the value.
    padded = resolve_default_value(
        seeded_db,
        country_of_origin=UNIQUE_COUNTRY,
        cn_code=ROUTED_CN,
        production_route="  (a)  ",
    )
    assert padded.status == RESOLUTION_RESOLVED
    assert padded.value is not None
    assert padded.value.id == with_route.value.id


def test_other_countries_group_is_never_a_fallback(seeded_db) -> None:
    ensure_platform_precursor_default_catalog(seeded_db)
    not_fallen_back = resolve_default_value(
        seeded_db,
        country_of_origin=UNIQUE_COUNTRY,
        cn_code=GROUP_ONLY_CN,
        production_route=GROUP_ONLY_ROUTE,
    )
    assert not_fallen_back.status == RESOLUTION_UNRESOLVED
    assert "never selected automatically" in not_fallen_back.other_countries_note

    explicit = resolve_default_value(
        seeded_db,
        country_of_origin=OTHER_COUNTRIES_GROUP_NAME,
        cn_code=GROUP_ONLY_CN,
        production_route=GROUP_ONLY_ROUTE,
    )
    assert explicit.status == RESOLUTION_RESOLVED
    assert explicit.value is not None
    assert explicit.value.is_other_countries_group is True


def test_country_matching_is_case_insensitive(seeded_db) -> None:
    ensure_platform_precursor_default_catalog(seeded_db)
    resolution = resolve_default_value(
        seeded_db, country_of_origin="  aLBANia ", cn_code=UNIQUE_CN_NORMALIZED
    )
    assert resolution.status == RESOLUTION_RESOLVED
    assert resolution.value is not None
    assert resolution.value.country_name == UNIQUE_COUNTRY


# --------------------------------------------------------------------------------------
# Snapshots
# --------------------------------------------------------------------------------------


def test_snapshot_carries_dataset_and_value_provenance(seeded_db) -> None:
    dataset = ensure_platform_precursor_default_catalog(seeded_db)
    resolution = resolve_default_value(
        seeded_db, country_of_origin=UNIQUE_COUNTRY, cn_code=UNIQUE_CN_NORMALIZED
    )
    assert resolution.value is not None
    value = seeded_db.get(CbamPrecursorDefaultValue, resolution.value.id)
    assert value is not None

    snapshot = build_default_snapshot(
        dataset=dataset,
        value=value,
        lookup_key=resolution.lookup_key,
        requested_country=UNIQUE_COUNTRY,
        requested_cn=UNIQUE_CN_NORMALIZED,
        requested_route=None,
        requested_description=None,
    )
    assert snapshot["snapshotVersion"] == 1
    assert snapshot["dataset"]["contentChecksum"] == EXPECTED_CONTENT_CHECKSUM
    assert snapshot["dataset"]["sourceWorkbookSha256"] == DV_WORKBOOK_SHA256
    assert snapshot["value"]["cnNormalizedCode"] == UNIQUE_CN_NORMALIZED
    assert snapshot["requestedKeys"]["lookupKey"] == resolution.lookup_key
    assert snapshot["units"]["specificDirect"] == SPECIFIC_DIRECT_UNIT

    direct, indirect = snapshot_specific_values(snapshot)
    assert direct == UNIQUE_DIRECT
    assert indirect == UNIQUE_INDIRECT


def test_snapshot_specific_values_ignore_non_numeric_states() -> None:
    assert snapshot_specific_values(None) == (None, None)
    assert snapshot_specific_values({}) == (None, None)
    dashed = {
        "value": {
            "directValue": None,
            "directValueStatus": DV_STATUS_DASH,
            "indirectValue": "0.03",
            "indirectValueStatus": DV_STATUS_NUMERIC,
        }
    }
    assert snapshot_specific_values(dashed) == (None, Decimal("0.03"))
