"""Phase 12A Official SEE export final acceptance (DB golden + LibreOffice + API)."""

from __future__ import annotations

import json
import shutil
import threading
import uuid
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from tests.cbam_official_see_golden import (
    CN_NUTS,
    CN_SCREWS,
    seed_official_see_golden,
)
from tests.helpers import api_login, auth_headers, current_org_id

from ecotrace.core.exceptions import BusinessRuleError
from ecotrace.modules.cbam.application.official_see_export.constants import (
    CODE_FORMULA_PARITY_FAILED,
    XLSX_MIME,
)
from ecotrace.modules.cbam.application.official_see_export.leakage import scan_example_leakage
from ecotrace.modules.cbam.application.official_see_export.mapping.loader import load_manifest
from ecotrace.modules.cbam.application.official_see_export.package_inventory import (
    assert_package_acceptable_after_libreoffice,
    inventory_package,
)
from ecotrace.modules.cbam.application.official_see_export.parity import (
    SEE_IJK_DECIMAL_PLACES,
    SEE_IJK_NUMBER_FORMAT,
    assert_no_formula_errors_or_raise,
    compare_outputs,
    scan_output_formula_errors,
    workbook_places_for_cell,
)
from ecotrace.modules.cbam.application.official_see_export.recalc import (
    resolve_soffice_path,
    soffice_available,
)
from ecotrace.modules.cbam.application.official_see_export.schemas import (
    OfficialSeeExportCreateRequest,
)
from ecotrace.modules.cbam.application.official_see_export.service import (
    create_official_see_export_run,
    download_official_see_export_artifact,
    list_official_see_export_artifacts,
)
from ecotrace.modules.cbam.application.official_see_export.writer import build_used_input_map
from ecotrace.modules.cbam.application.product_embedded_emissions_service import (
    get_product_embedded_emissions_result,
)
from ecotrace.modules.cbam.infrastructure.models import (
    CbamOfficialSeeExportArtifact,
    CbamOfficialSeeExportRun,
)

GOLDEN_DIR = Path("/tmp/ecotrace-see-golden-acceptance")


def _quantize_3(value: Decimal) -> Decimal:
    assert workbook_places_for_cell(SEE_IJK_NUMBER_FORMAT) == SEE_IJK_DECIMAL_PLACES == 3
    return value.quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)


def _pee_specifics(db: Session, scenario) -> list[dict[str, Decimal | str]]:
    detail = get_product_embedded_emissions_result(
        db,
        scenario.user,
        scenario.organization.id,
        scenario.binding.id,
        scenario.pee_result_id,
    )
    rows = []
    for product in detail.products:
        rows.append(
            {
                "cn": str(product.get("cnNormalizedCode") or ""),
                "direct": Decimal(str(product["specificDirect"])),
                "indirect": Decimal(str(product["specificIndirect"])),
                "total": Decimal(str(product["specificTotal"])),
                "name": str(product.get("productName") or ""),
            }
        )
    return rows


def _cell_dec(wb, sheet: str, cell: str) -> Decimal:
    raw = wb[sheet][cell].value
    assert raw is not None and not (isinstance(raw, str) and raw.startswith("#")), (
        f"{sheet}!{cell}={raw!r}"
    )
    return Decimal(str(raw))


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice soffice required")
def test_parity_acceptance_db_golden(seeded_db: Session) -> None:
    """STOP condition: PEE V2 snapshot I/J/K must match workbook at format 0.000 (3 dp)."""
    assert resolve_soffice_path() is not None
    scenario = seed_official_see_golden(seeded_db)
    pee_rows = _pee_specifics(seeded_db, scenario)
    assert len(pee_rows) == 2

    run = create_official_see_export_run(
        seeded_db,
        scenario.user,
        scenario.organization.id,
        scenario.binding.id,
        OfficialSeeExportCreateRequest(client_request_id=uuid.uuid4()),
    )
    assert run.generation_status == "COMPLETED"
    assert run.formula_parity_status == "PASSED"
    assert run.output_sha256

    artifacts = list_official_see_export_artifacts(
        seeded_db, scenario.user, scenario.organization.id, run.id
    )
    assert len(artifacts) == 1
    artifact, path = download_official_see_export_artifact(
        seeded_db, scenario.user, scenario.organization.id, artifacts[0].id
    )
    assert path.is_file()

    GOLDEN_DIR.mkdir(parents=True, exist_ok=True)
    golden_xlsx = GOLDEN_DIR / artifact.file_name
    shutil.copy2(path, golden_xlsx)

    wb = load_workbook(path, data_only=True)
    wb_fmt = load_workbook(path, data_only=False)
    mismatch_table: list[dict[str, str]] = []
    cn_codes = []
    pee_by_cn = {str(r["cn"]): r for r in pee_rows}
    for i in range(len(pee_rows)):
        row = 10 + i
        comm = 26 + i
        f_val = wb["Summary_Products"][f"F{row}"].value
        cn = str(f_val) if f_val is not None else ""
        cn_codes.append(cn)
        assert cn in {CN_SCREWS, CN_NUTS}
        pee = pee_by_cn.get(cn)
        assert pee is not None, f"No PEE row for CN {cn}"
        for col, key in (("I", "direct"), ("J", "indirect"), ("K", "total")):
            assert wb_fmt["Summary_Products"][f"{col}{row}"].number_format == SEE_IJK_NUMBER_FORMAT
            expected = _quantize_3(pee[key])  # type: ignore[arg-type]
            actual_sp = _quantize_3(_cell_dec(wb, "Summary_Products", f"{col}{row}"))
            actual_sc = _quantize_3(_cell_dec(wb, "Summary_Communication", f"{col}{comm}"))
            if actual_sp != expected or actual_sc != expected:
                mismatch_table.append(
                    {
                        "product": str(i),
                        "cn": cn,
                        "cell": f"{col}{row}/{col}{comm}",
                        "pee": str(pee[key]),
                        "pee_3dp": str(expected),
                        "summary_products": str(actual_sp),
                        "summary_communication": str(actual_sc),
                    }
                )
        g_sp = wb["Summary_Products"][f"G{row}"].value
        g_sc = wb["Summary_Communication"][f"G{comm}"].value
        assert g_sp is not None and not str(g_sp).startswith("#")
        assert g_sc is not None and not str(g_sc).startswith("#")

    report = {
        "peeResultId": str(scenario.pee_result_id),
        "outputSha256": run.output_sha256,
        "seeNumberFormat": SEE_IJK_NUMBER_FORMAT,
        "decimalPlaces": SEE_IJK_DECIMAL_PLACES,
        "cnCodes": cn_codes,
        "peeRows": [
            {
                "cn": r["cn"],
                "specificDirect": str(r["direct"]),
                "specificIndirect": str(r["indirect"]),
                "specificTotal": str(r["total"]),
                "direct3dp": str(_quantize_3(r["direct"])),  # type: ignore[arg-type]
                "indirect3dp": str(_quantize_3(r["indirect"])),  # type: ignore[arg-type]
                "total3dp": str(_quantize_3(r["total"])),  # type: ignore[arg-type]
            }
            for r in pee_rows
        ],
        "mismatches": mismatch_table,
    }
    (GOLDEN_DIR / "report.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    if mismatch_table:
        pytest.fail(
            "Phase 12A STOP: PEE↔workbook I/J/K parity failed at 3 dp (format 0.000):\n"
            + json.dumps(mismatch_table, indent=2)
        )

    manifest = load_manifest()
    assert scan_output_formula_errors(path, manifest, product_count=2) == []
    assert_no_formula_errors_or_raise(path, manifest, product_count=2)

    from ecotrace.modules.cbam.application.official_see_export.context import (
        load_official_see_context,
    )

    ctx = load_official_see_context(
        seeded_db, scenario.user, scenario.organization.id, scenario.binding.id
    )
    parity = compare_outputs(path, manifest, ctx.expected_outputs)
    assert parity == [], parity[:5]
    used = set(build_used_input_map(ctx, manifest).keys())
    assert scan_example_leakage(path, manifest, used_input_keys=used) == []


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice soffice required")
def test_api_acceptance_on_golden(client: TestClient, seeded_db: Session) -> None:
    scenario = seed_official_see_golden(seeded_db)
    seeded_db.commit()

    token = api_login(client, "orgadmin@ecotrace.dev", "EcoTraceOrgAdmin!2024")
    org_id = current_org_id(client, token)
    assert org_id == str(scenario.organization.id)
    binding_id = str(scenario.binding.id)
    base = (
        f"/api/v1/cbam/organizations/{org_id}/reporting-period-bindings/"
        f"{binding_id}/official-see-export"
    )
    headers = auth_headers(token)

    readiness = client.get(f"{base}/readiness", headers=headers)
    assert readiness.status_code == 200, readiness.text
    body = readiness.json()
    assert body["ready"] is True, body
    assert body["sofficeAvailable"] is True

    client_request_id = str(uuid.uuid4())
    first = client.post(
        f"{base}/executions",
        headers=headers,
        json={"clientRequestId": client_request_id},
    )
    assert first.status_code == 201, first.text
    first_json = first.json()
    assert first_json["generationStatus"] == "COMPLETED"
    assert first_json["formulaParityStatus"] == "PASSED"
    sha = first_json["outputSha256"]
    run_id = first_json["id"]

    replay = client.post(
        f"{base}/executions",
        headers=headers,
        json={"clientRequestId": client_request_id},
    )
    assert replay.status_code == 200, replay.text
    assert replay.json()["idempotentReplay"] is True
    assert replay.json()["outputSha256"] == sha
    assert replay.json()["id"] == run_id

    # Same key, different fingerprint → 409
    run_row = seeded_db.get(CbamOfficialSeeExportRun, uuid.UUID(run_id))
    assert run_row is not None
    run_row.source_fingerprint = "0" * 64
    seeded_db.commit()
    reused = client.post(
        f"{base}/executions",
        headers=headers,
        json={"clientRequestId": client_request_id},
    )
    assert reused.status_code == 409, reused.text

    # Concurrent same key → single artifact sha
    concurrent_key = str(uuid.uuid4())
    results: list[tuple[int, str | None]] = []

    def _post() -> None:
        r = client.post(
            f"{base}/executions",
            headers=headers,
            json={"clientRequestId": concurrent_key},
        )
        sha_out = r.json().get("outputSha256") if r.status_code in {200, 201} else None
        results.append((r.status_code, sha_out))

    t1 = threading.Thread(target=_post)
    t2 = threading.Thread(target=_post)
    t1.start()
    t2.start()
    t1.join()
    t2.join()
    assert all(code in {200, 201, 409} for code, _ in results)
    completed_shas = {out for code, out in results if code in {200, 201} and out}
    assert len(completed_shas) == 1

    artifacts = client.get(
        f"/api/v1/cbam/organizations/{org_id}/official-see-export/runs/{run_id}/artifacts",
        headers=headers,
    )
    assert artifacts.status_code == 200
    art = artifacts.json()[0]
    assert "storageUri" not in art
    assert art["fileName"].endswith(".xlsx")
    assert "/" not in art["fileName"]
    assert ".." not in art["fileName"]

    download = client.get(
        f"/api/v1/cbam/organizations/{org_id}/official-see-export/artifacts/{art['id']}/download",
        headers=headers,
    )
    assert download.status_code == 200
    assert XLSX_MIME.split(";")[0] in (download.headers.get("content-type") or "")
    assert download.content[:2] == b"PK"

    viewer = api_login(client, "viewer@ecotrace.dev", "EcoTraceViewer!2024")
    v_headers = auth_headers(viewer)
    assert client.get(f"{base}/readiness", headers=v_headers).status_code == 200
    denied = client.post(
        f"{base}/executions",
        headers=v_headers,
        json={"clientRequestId": str(uuid.uuid4())},
    )
    assert denied.status_code in {403, 401}

    foreign = str(uuid.uuid4())
    cross = client.get(
        f"/api/v1/cbam/organizations/{foreign}/reporting-period-bindings/"
        f"{binding_id}/official-see-export/readiness",
        headers=headers,
    )
    assert cross.status_code in {403, 404}

    # Failed parity → no COMPLETED downloadable artifact for that run
    fail_key = uuid.uuid4()
    with patch(
        "ecotrace.modules.cbam.application.official_see_export.service.assert_parity_or_raise",
        side_effect=BusinessRuleError(
            "forced parity fail",
            code=CODE_FORMULA_PARITY_FAILED,
            details=[{"code": CODE_FORMULA_PARITY_FAILED}],
        ),
    ):
        failed = client.post(
            f"{base}/executions",
            headers=headers,
            json={"clientRequestId": str(fail_key)},
        )
    assert failed.status_code == 400
    failed_run = seeded_db.execute(
        select(CbamOfficialSeeExportRun).where(
            CbamOfficialSeeExportRun.organization_id == scenario.organization.id,
            CbamOfficialSeeExportRun.client_request_id == fail_key,
        )
    ).scalar_one()
    assert failed_run.generation_status == "FAILED"
    arts = list(
        seeded_db.execute(
            select(CbamOfficialSeeExportArtifact).where(
                CbamOfficialSeeExportArtifact.export_run_id == failed_run.id
            )
        )
        .scalars()
        .all()
    )
    assert arts == []


@pytest.mark.skipif(not soffice_available(), reason="LibreOffice soffice required")
def test_package_allowlist_after_lo_on_golden_write(seeded_db: Session, tmp_path: Path) -> None:
    from ecotrace.modules.cbam.application.official_see_export.context import (
        load_official_see_context,
    )
    from ecotrace.modules.cbam.application.official_see_export.recalc import (
        cleanup_recalc_dir,
        recalculate_workbook,
    )
    from ecotrace.modules.cbam.application.official_see_export.service import (
        ensure_official_see_template,
    )
    from ecotrace.modules.cbam.application.official_see_export.writer import (
        write_official_see_workbook,
    )

    scenario = seed_official_see_golden(seeded_db)
    ctx = load_official_see_context(
        seeded_db, scenario.user, scenario.organization.id, scenario.binding.id
    )
    template = ensure_official_see_template()
    staging = tmp_path / "staging.xlsx"
    write_official_see_workbook(template_path=template, output_path=staging, ctx=ctx)
    before = inventory_package(staging)
    recalc = recalculate_workbook(staging, timeout_seconds=300)
    try:
        after = inventory_package(recalc)
        assert_package_acceptable_after_libreoffice(before, after)
    finally:
        cleanup_recalc_dir(recalc)
