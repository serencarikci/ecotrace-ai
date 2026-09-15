#!/usr/bin/env python3
"""Deterministic DV workbook → normalized precursor default-value catalog seed JSON.

Does not commit the binary workbook. Point WORKBOOK_PATH at the local DV file.

Usage:
  WORKBOOK_PATH=/path/to/DVs_as_adopted_v20260204.xlsx \\
    python scripts/extract_cbam_precursor_default_catalog.py

Output:
  src/ecotrace/modules/cbam/data/cbam_precursor_defaults_v20260204.json
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = ROOT.parent
DEFAULT_OUT = (
    ROOT / 'src/ecotrace/modules/cbam/data/cbam_precursor_defaults_v20260204.json'
)
DEFAULT_WORKBOOK = REPO_ROOT / 'local-reference' / 'DVs_as_adopted_v20260204.xlsx'

DATASET_CODE = 'CBAM_EU_DEFAULT_VALUES'
DATASET_VERSION = 'IR_2025_2621_v20260204'
SOURCE_TEMPLATE_VERSION = '1'
VALID_FROM = '2026-02-04'
REGULATION_REF = 'Commission Implementing Regulation (EU) 2025/2621'

META_SHEETS = {'Overview', 'Version History'}
OTHER_SHEET = '_Other Countries and Territorie'

HEADER_MARKERS = (
    'Product CN Code',
    'Description',
    'Default Value',
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _normalize_cn(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    if not text:
        return None
    digits = re.sub(r'[^0-9]', '', text)
    return digits or None


def _display_cn(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _route_value(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).replace('\xa0', ' ').strip()
    if not text or text.upper() in {'N/A', 'NA', '-', '–', '_', '—'}:
        return None
    return text


def _parse_numeric_or_status(value: object) -> tuple[str | None, str]:
    """Return (decimal_string_or_None, status)."""
    if value is None:
        return None, 'BLANK'
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return str(Decimal(str(value))), 'NUMERIC'
        except InvalidOperation:
            return None, 'INVALID'
    text = str(value).replace('\xa0', ' ').strip()
    if not text:
        return None, 'BLANK'
    lowered = text.lower()
    if lowered in {'n/a', 'na'}:
        return None, 'NA'
    if lowered in {'see below'}:
        return None, 'SEE_BELOW'
    if text in {'-', '–', '_', '—'}:
        return None, 'DASH'
    try:
        return str(Decimal(text.replace(',', ''))), 'NUMERIC'
    except InvalidOperation:
        return None, 'TEXT'


def _is_category_row(cn_cell: object, description: object) -> bool:
    cn_text = '' if cn_cell is None else str(cn_cell).strip()
    if cn_text:
        return False
    desc = '' if description is None else str(description).replace('\xa0', ' ').strip()
    return bool(desc)


def extract(workbook_path: Path) -> dict:
    digest = _sha256(workbook_path)
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        sheet_names = list(wb.sheetnames)
        version_notes = None
        version_date = None
        if 'Version History' in wb.sheetnames:
            ws = wb['Version History']
            rows = list(ws.iter_rows(min_row=3, max_row=5, max_col=3, values_only=True))
            if rows and rows[0][0] is not None:
                version_notes = str(rows[0][2]) if rows[0][2] is not None else None
                version_date = str(rows[0][1]) if rows[0][1] is not None else None

        values: list[dict] = []
        for sheet_name in sheet_names:
            if sheet_name in META_SHEETS:
                continue
            ws = wb[sheet_name]
            country_name = None
            category = None
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=1, max_col=9, values_only=True), start=1
            ):
                a, b, c, d, e, f, g, h, i = (list(row) + [None] * 9)[:9]
                if row_idx == 1:
                    country_name = (
                        'Other Countries and Territories'
                        if sheet_name == OTHER_SHEET
                        else (str(a).strip() if a is not None else sheet_name)
                    )
                    continue
                if row_idx == 2:
                    continue
                if _is_category_row(a, b):
                    category = str(b).replace('\xa0', ' ').strip()
                    continue
                cn_norm = _normalize_cn(a)
                if cn_norm is None:
                    continue
                direct_val, direct_status = _parse_numeric_or_status(c)
                indirect_val, indirect_status = _parse_numeric_or_status(d)
                total_val, total_status = _parse_numeric_or_status(e)
                marked_2026, _ = _parse_numeric_or_status(f)
                marked_2027, _ = _parse_numeric_or_status(g)
                marked_2028, _ = _parse_numeric_or_status(h)
                route = _route_value(i)
                description = (
                    str(b).replace('\xa0', ' ').strip() if b is not None else None
                ) or None
                values.append(
                    {
                        'countryName': country_name,
                        'sourceSheet': sheet_name,
                        'sourceRow': row_idx,
                        'isOtherCountriesGroup': sheet_name == OTHER_SHEET,
                        'cnNormalizedCode': cn_norm,
                        'cnDisplayCode': _display_cn(a),
                        'goodsCategory': category,
                        'goodsDescription': description,
                        'productionRoute': route,
                        'directValue': direct_val,
                        'directValueStatus': direct_status,
                        'indirectValue': indirect_val,
                        'indirectValueStatus': indirect_status,
                        'totalValue': total_val,
                        'totalValueStatus': total_status,
                        'directUnit': None,
                        'indirectUnit': None,
                        'totalUnit': None,
                        'unitNote': (
                            'DV workbook does not label emission units; '
                            'SEE applies specific SEE as tCO2e per goods unit.'
                        ),
                        'markedUpTotals': {
                            '2026': marked_2026,
                            '2027': marked_2027,
                            '2028_and_onwards': marked_2028,
                        },
                        'originalKeys': {
                            'sheetName': sheet_name,
                            'countryCell': country_name,
                            'cnCodeCell': _display_cn(a),
                            'descriptionCell': description,
                            'productionRouteCell': (
                                None
                                if i is None
                                else str(i).replace('\xa0', ' ').strip() or None
                            ),
                        },
                    }
                )
    finally:
        wb.close()

    values.sort(
        key=lambda r: (
            r['countryName'] or '',
            r['cnNormalizedCode'] or '',
            r['productionRoute'] or '',
            r['goodsDescription'] or '',
            r['sourceRow'],
        )
    )

    payload = {
        'dataset': {
            'datasetCode': DATASET_CODE,
            'datasetVersion': DATASET_VERSION,
            'sourceWorkbookName': workbook_path.name,
            'sourceWorkbookSha256': digest,
            'sourceTemplateVersion': SOURCE_TEMPLATE_VERSION,
            'regulationReference': REGULATION_REF,
            'versionHistoryDate': version_date,
            'versionHistoryNotes': version_notes,
            'validFrom': VALID_FROM,
            'validUntil': None,
            'unitLabelInWorkbook': None,
            'seeApplicationUnitDirect': 'tCO2e/t',
            'seeApplicationUnitIndirect': 'tCO2e/t',
            'resolverKey': [
                'countryName',
                'cnNormalizedCode',
                'productionRoute',
                'goodsDescription',
            ],
            'ambiguousWithoutDescription': True,
            'valueCount': len(values),
        },
        'values': values,
    }
    checksum = hashlib.sha256(
        json.dumps(
            {**payload, 'dataset': {k: v for k, v in payload['dataset'].items() if k != 'contentChecksum'}},
            ensure_ascii=False,
            sort_keys=True,
            separators=(',', ':'),
        ).encode()
    ).hexdigest()
    payload['dataset']['contentChecksum'] = checksum
    return payload


def main() -> int:
    workbook = Path(os.environ.get('WORKBOOK_PATH', str(DEFAULT_WORKBOOK)))
    out = Path(os.environ.get('OUT_PATH', str(DEFAULT_OUT)))
    if not workbook.is_file():
        print(f'Workbook not found: {workbook}', file=sys.stderr)
        return 1
    payload = extract(workbook)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8'
    )
    print(f'Wrote {out}')
    print(f'values={payload["dataset"]["valueCount"]}')
    print(f'checksum={payload["dataset"]["contentChecksum"]}')
    print(f'sha256={payload["dataset"]["sourceWorkbookSha256"]}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
